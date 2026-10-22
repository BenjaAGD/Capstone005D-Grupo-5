
import os
import time
from datetime import datetime
from fastapi import APIRouter, Body
from fastapi.responses import FileResponse, JSONResponse, Response

from modulos_base import (
    SEVERIDADES, agentes_analizables, nombre_corto, es_integrador, tiene_info, score_de,
    etiqueta_score, banda_score, senales_de, lista_textos, id_estable, info_documento,
    cargar_analisis, similitud_titulos, esc_pdf, html_a_pdf, pagina_pdf,
    seccion, parrafo, nota_pdf, tiles_pdf, tabla_pdf, color_texto, COLOR_SEV, NAVY,
)

UMBRAL_MISMA_SENAL = 0.55


def _conteo_vacio():
    return {s: 0 for s in SEVERIDADES}


def construir_dashboard(sb, document_id):
    doc = info_documento(sb, document_id)
    if doc is None:
        return None
    analisis = cargar_analisis(sb, document_id)

    agentes, senales, omisiones, preguntas, recomendaciones = [], [], [], [], []
    for key in agentes_analizables():
        res = analisis.get(key)
        base = {"key": key, "nombre": nombre_corto(key), "transversal": es_integrador(key)}
        if res is None:
            agentes.append({**base, "estado": "sin_analizar", "con_info": False, "score": None,
                            "etiqueta": "Sin analizar", "banda": "sin_analizar",
                            "senales": _conteo_vacio(), "n_omisiones": 0, "n_preguntas": 0,
                            "n_recomendaciones": 0, "resumen": ""})
            continue
        s_list = senales_de(key, res)
        cont = _conteo_vacio()
        for s in s_list:
            cont[s["severidad"]] += 1
        om = lista_textos(res.get("omisiones"), "descripcion", "texto")
        pr = lista_textos(res.get("preguntas"), "pregunta", "texto")
        rc = lista_textos(res.get("recomendaciones"), "descripcion", "recomendacion", "texto")
        con_info = tiene_info(res)
        sc = score_de(res)
        agentes.append({**base, "estado": "ok", "con_info": con_info, "score": sc,
                        "etiqueta": etiqueta_score(sc) if con_info else "Sin información relevante",
                        "banda": banda_score(sc) if con_info else "sin_info",
                        "senales": cont, "n_omisiones": len(om), "n_preguntas": len(pr),
                        "n_recomendaciones": len(rc), "resumen": str(res.get("resumen_ejecutivo") or "")})
        senales.extend(s_list)
        omisiones += [{"agente": key, "agente_nombre": nombre_corto(key), "texto": t} for t in om]
        preguntas += [{"agente": key, "agente_nombre": nombre_corto(key), "texto": t} for t in pr]
        recomendaciones += [{"id": id_estable(document_id, key, t), "agente": key,
                             "agente_nombre": nombre_corto(key), "texto": t} for t in rc]

    con_score = [a["score"] for a in agentes if a["con_info"]]
    global_score = round(sum(con_score) / len(con_score)) if con_score else None
    total_sen = _conteo_vacio()
    for s in senales:
        total_sen[s["severidad"]] += 1
    orden = {sv: i for i, sv in enumerate(SEVERIDADES)}
    senales.sort(key=lambda s: (orden[s["severidad"]], s["transversal"]))

    return {
        "ok": True,
        "documento": doc,
        "resumen": {
            "score_global": global_score,
            "etiqueta": etiqueta_score(global_score) if global_score is not None else "Sin datos",
            "banda": banda_score(global_score) if global_score is not None else "sin_datos",
            "agentes_total": len(agentes),
            "agentes_analizados": sum(1 for a in agentes if a["estado"] == "ok"),
            "agentes_con_info": len(con_score),
            "senales": total_sen,
            "senales_total": len(senales),
            "senales_transversales": sum(1 for s in senales if s["transversal"]),
            "omisiones": len(omisiones),
            "preguntas": len(preguntas),
            "recomendaciones": len(recomendaciones),
        },
        "agentes": agentes,
        "senales": senales,
        "omisiones": omisiones,
        "preguntas": preguntas,
        "recomendaciones": recomendaciones,
    }


_CACHE_DASH = {}
CACHE_DASH_SEG = 45


def dashboard_cacheado(sb, document_id, refrescar=False):
    """Igual que construir_dashboard, pero reutiliza el resultado de los últimos 45 s (el análisis no cambia a cada rato)."""
    ahora = time.time()
    hit = _CACHE_DASH.get(document_id)
    if hit and not refrescar and ahora - hit[0] < CACHE_DASH_SEG:
        return hit[1]
    data = construir_dashboard(sb, document_id)
    if data is not None:
        _CACHE_DASH[document_id] = (ahora, data)
    return data


def comparar_documentos(sb, id_a, id_b):
    """A = documento base (anterior), B = documento nuevo. Todo sale de lo ya guardado."""
    da, db_ = dashboard_cacheado(sb, id_a), dashboard_cacheado(sb, id_b)
    if da is None or db_ is None:
        return None
    por_agente = []
    for aa in da["agentes"]:
        bb = next(x for x in db_["agentes"] if x["key"] == aa["key"])
        delta = (bb["score"] - aa["score"]) if (aa["score"] is not None and bb["score"] is not None
                                                and aa["con_info"] and bb["con_info"]) else None
        por_agente.append({"key": aa["key"], "nombre": aa["nombre"], "transversal": aa["transversal"],
                           "score_a": aa["score"] if aa["con_info"] else None,
                           "score_b": bb["score"] if bb["con_info"] else None,
                           "delta": delta,
                           "senales_a": sum(aa["senales"].values()), "senales_b": sum(bb["senales"].values())})

    nuevas, resueltas, persisten = [], [], []
    usadas_b = set()
    for sa in da["senales"]:
        match = None
        for j, sb_ in enumerate(db_["senales"]):
            if j in usadas_b or sb_["agente"] != sa["agente"]:
                continue
            if similitud_titulos(sa["titulo"], sb_["titulo"]) >= UMBRAL_MISMA_SENAL:
                match = j
                break
        if match is None:
            resueltas.append(sa)
        else:
            usadas_b.add(match)
            persisten.append({**db_["senales"][match], "severidad_anterior": sa["severidad"]})
    nuevas = [s for j, s in enumerate(db_["senales"]) if j not in usadas_b]
    ga, gb = da["resumen"]["score_global"], db_["resumen"]["score_global"]
    return {
        "ok": True,
        "a": da["documento"], "b": db_["documento"],
        "global": {"a": ga, "b": gb, "delta": (gb - ga) if (ga is not None and gb is not None) else None},
        "senales_por_severidad": {"a": da["resumen"]["senales"], "b": db_["resumen"]["senales"]},
        "agentes": por_agente,
        "nuevas": nuevas, "resueltas": resueltas, "persisten": persisten,
        "nota": "El emparejamiento de señales entre documentos es aproximado (similitud de títulos).",
    }


_ESTADOS_PLAN = {"pendiente": "Pendiente", "curso": "En curso", "hecho": "Hecha"}
_PLURAL_SEV = {"crítica": "críticas", "alta": "altas", "media": "medias", "baja": "bajas"}
_BANDAS_TXT = (("90 a 100", "Sin observaciones", "No hay hallazgos que requieran acción."),
               ("75 a 89", "Menor atención", "Observaciones menores; conviene revisarlas."),
               ("55 a 74", "Requiere revisión", "Hay temas que deben revisarse antes de la reunión."),
               ("40 a 54", "Observaciones serias", "Hay riesgos relevantes que exigen un plan de acción."),
               ("0 a 39", "Observaciones críticas", "Situación que requiere atención inmediata."))


def _sev_txt(sv):
    return color_texto(sv.upper(), COLOR_SEV.get(sv, NAVY))


def _fecha_hoy():
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def html_pdf_dashboard(d, acciones=None):
    """PDF corporativo del Dashboard. `acciones` = {id_recomendación: {estado, responsable, fecha}} (viene del navegador)."""
    acciones = acciones or {}
    r = d["resumen"]
    sv = r["senales"]
    doc = d["documento"]

    # 1. Resumen ejecutivo
    con_info = sorted((a for a in d["agentes"] if a["con_info"] and a["score"] is not None), key=lambda a: a["score"])
    if r["score_global"] is None:
        lectura = "Aún no hay análisis suficientes para calcular un puntaje global."
    else:
        lectura = (f"El puntaje global de preparación es {r['score_global']} sobre 100 ({r['etiqueta']}). "
                   f"Se detectaron {r['senales_total']} señales: {sv['crítica']} críticas, {sv['alta']} altas, "
                   f"{sv['media']} medias y {sv['baja']} bajas. ")
        if con_info:
            peores = ", ".join(f"{a['nombre']} ({a['score']})" for a in con_info[:3])
            lectura += f"Las áreas que requieren mayor atención son: {peores}. "
        lectura += (f"Además, los agentes identificaron {r['omisiones']} omisiones, {r['preguntas']} preguntas probables "
                    f"del directorio y {r['recomendaciones']} recomendaciones.")
    gcol = {"ok": "#2F6B4F", "leve": "#2F6B4F", "revision": "#A67C00", "serio": "#C2571A", "critico": "#8E1B2E"}.get(r["banda"], NAVY)
    tiles = tiles_pdf([
        (f"{r['score_global']}/100" if r["score_global"] is not None else "-", "Puntaje global", gcol),
        (sv["crítica"], "Señales críticas", COLOR_SEV["crítica"]), (sv["alta"], "Señales altas", COLOR_SEV["alta"]),
        (sv["media"], "Señales medias", COLOR_SEV["media"]), (sv["baja"], "Señales bajas", COLOR_SEV["baja"])])
    cuerpo = seccion(1, "Resumen ejecutivo") + parrafo(lectura) + tiles

    # 2. Semáforo por área
    orden = sorted(d["agentes"], key=lambda a: (a["score"] is None, a["score"] if a["score"] is not None else 999, a["nombre"]))
    filas = []
    for a in orden:
        s_ = a["senales"]
        puntaje = str(a["score"]) if a["score"] is not None and a.get("con_info") else "-"
        nombre = f"<b>{esc_pdf(a['nombre'])}</b>" + (" <span class='small muted'>(transversal)</span>" if a.get("transversal") else "")
        filas.append([nombre, puntaje, esc_pdf(a["etiqueta"]), str(s_["crítica"]), str(s_["alta"]), str(s_["media"]),
                      str(s_["baja"]), str(a["n_omisiones"]), str(a["n_preguntas"])])
    cuerpo += seccion(2, "Estado por área de análisis")
    cuerpo += parrafo("Cada área evalúa un aspecto distinto del documento. Se ordenan de menor a mayor puntaje para ver primero las más delicadas.")
    cuerpo += tabla_pdf(["Área", "Puntaje", "Estado", "Crít.", "Altas", "Medias", "Bajas", "Omis.", "Preg."], filas,
                        ["20%", "9%", "22%", "8%", "8%", "9%", "8%", "8%", "8%"], num_cols=(1, 3, 4, 5, 6, 7, 8))

    # 3. Señales
    orden_sev = {k: i for i, k in enumerate(SEVERIDADES)}
    sens = sorted(d["senales"], key=lambda x: (orden_sev.get(x["severidad"], 9), x.get("transversal", False)))
    cuerpo += seccion(3, f"Señales detectadas ({len(sens)})")
    if sens:
        filas = [[_sev_txt(x["severidad"]), esc_pdf(x["agente_nombre"]) + (" <span class='small muted'>(transv.)</span>" if x.get("transversal") else ""),
                  f"<b>{esc_pdf(x['titulo'])}</b><br/><span class='muted'>{esc_pdf(x.get('descripcion', ''))}</span>"] for x in sens]
        cuerpo += tabla_pdf(["Severidad", "Área", "Señal"], filas, ["12%", "18%", "70%"])
    else:
        cuerpo += parrafo("No se detectaron señales.")

    # 4. Plan de acción
    hechas = 0
    filas = []
    for rec in d["recomendaciones"]:
        a = acciones.get(rec["id"]) or {}
        est = a.get("estado") if a.get("estado") in _ESTADOS_PLAN else "pendiente"
        hechas += est == "hecho"
        filas.append([esc_pdf(rec["agente_nombre"]), esc_pdf(rec["texto"]), esc_pdf(a.get("responsable", "") or "-"),
                      esc_pdf(a.get("fecha", "") or "-"), _ESTADOS_PLAN[est]])
    cuerpo += seccion(4, f"Plan de acción ({hechas} de {len(d['recomendaciones'])} acciones hechas)")
    cuerpo += (tabla_pdf(["Área", "Acción recomendada", "Responsable", "Fecha límite", "Estado"], filas, ["14%", "46%", "16%", "12%", "12%"])
               if filas else parrafo("No hay recomendaciones."))

    # 5. Cómo leer el puntaje
    cuerpo += seccion(5, "Cómo leer el puntaje")
    cuerpo += parrafo("El puntaje de cada área va de 0 a 100 y resume qué tan preparado está el documento en ese aspecto. "
                      "El puntaje global es el promedio de las áreas que encontraron información relevante.")
    cuerpo += tabla_pdf(["Puntaje", "Estado", "Significado"], [[b[0], f"<b>{esc_pdf(b[1])}</b>", esc_pdf(b[2])] for b in _BANDAS_TXT],
                        ["14%", "26%", "60%"])
    meta = [("Documento", doc["nombre"]), ("Fecha de emisión", _fecha_hoy()),
            ("Áreas analizadas", f"{r['agentes_analizados']} de {r['agentes_total']}"), ("Clasificación", "Confidencial")]
    return pagina_pdf("Dashboard de Preparación", "Indicadores para la Gerencia General", cuerpo, meta, "Dashboard de Preparación")


def _delta_txt(x):
    if x is None:
        return "-"
    return "0" if x == 0 else (f"+{x}" if x > 0 else str(x))


def html_pdf_comparacion(c):
    """PDF corporativo de la comparación entre dos documentos (A = anterior, B = actual)."""
    g = c["global"]
    sa, sb_ = c["senales_por_severidad"]["a"], c["senales_por_severidad"]["b"]
    nom_a, nom_b = c["a"]["nombre"], c["b"]["nombre"]
    if g["delta"] is None:
        lectura = "No hay puntaje global comparable entre ambos documentos."
    elif g["delta"] > 0:
        lectura = f"El puntaje global mejoró {g['delta']} puntos, de {g['a']} a {g['b']}."
    elif g["delta"] < 0:
        lectura = f"El puntaje global bajó {abs(g['delta'])} puntos, de {g['a']} a {g['b']}."
    else:
        lectura = f"El puntaje global se mantiene en {g['b']}."
    lectura += (f" Respecto del documento anterior hay {len(c['nuevas'])} señales nuevas, {len(c['resueltas'])} resueltas "
                f"y {len(c['persisten'])} que persisten.")
    dcol = NAVY if g["delta"] in (None, 0) else ("#2F6B4F" if g["delta"] > 0 else "#8E1B2E")
    cuerpo = seccion(1, "Resumen ejecutivo") + parrafo(lectura)
    cuerpo += tiles_pdf([
        (g["a"] if g["a"] is not None else "-", "Puntaje anterior", None),
        (g["b"] if g["b"] is not None else "-", "Puntaje actual", None),
        (_delta_txt(g["delta"]), "Variación", dcol),
        (len(c["nuevas"]), "Señales nuevas", COLOR_SEV["alta"]),
        (len(c["resueltas"]), "Señales resueltas", COLOR_SEV["baja"]),
        (len(c["persisten"]), "Señales que persisten", COLOR_SEV["media"])])

    cuerpo += seccion(2, "Señales por severidad")
    filas = [[_sev_txt(k), str(sa.get(k, 0)), str(sb_.get(k, 0)), _delta_txt(sb_.get(k, 0) - sa.get(k, 0))] for k in SEVERIDADES]
    cuerpo += tabla_pdf(["Severidad", "Anterior", "Actual", "Variación"], filas, ["34%", "22%", "22%", "22%"], num_cols=(1, 2, 3))

    cuerpo += seccion(3, "Variación del puntaje por área")
    filas = []
    for a in sorted(c["agentes"], key=lambda x: (x["delta"] is None, x["delta"] if x["delta"] is not None else 0)):
        dcol_a = "#8E1B2E" if (a["delta"] or 0) < 0 else ("#2F6B4F" if (a["delta"] or 0) > 0 else "#5F6877")
        filas.append([f"<b>{esc_pdf(a['nombre'])}</b>" + (" <span class='small muted'>(transversal)</span>" if a["transversal"] else ""),
                      str(a["score_a"]) if a["score_a"] is not None else "-", str(a["score_b"]) if a["score_b"] is not None else "-",
                      color_texto(_delta_txt(a["delta"]), dcol_a), f"{a['senales_a']} -> {a['senales_b']}"])
    cuerpo += tabla_pdf(["Área", "Anterior", "Actual", "Variación", "Señales"], filas, ["34%", "15%", "15%", "16%", "20%"], num_cols=(1, 2, 3, 4))

    def tabla_senales(lista, vacio, con_ant=False):
        if not lista:
            return parrafo(vacio)
        filas = []
        for x in lista:
            ant = ""
            if con_ant and x.get("severidad_anterior") and x["severidad_anterior"] != x["severidad"]:
                ant = f"<br/><span class='small muted'>Antes: {esc_pdf(x['severidad_anterior'])}</span>"
            filas.append([_sev_txt(x["severidad"]) + ant, esc_pdf(x["agente_nombre"]),
                          f"<b>{esc_pdf(x['titulo'])}</b><br/><span class='muted'>{esc_pdf(x.get('descripcion', ''))}</span>"])
        return tabla_pdf(["Severidad", "Área", "Señal"], filas, ["13%", "18%", "69%"])
    cuerpo += seccion(4, f"Señales nuevas ({len(c['nuevas'])})") + tabla_senales(c["nuevas"], "No aparecieron señales nuevas.")
    cuerpo += seccion(5, f"Señales resueltas ({len(c['resueltas'])})") + tabla_senales(c["resueltas"], "No hay señales resueltas.")
    cuerpo += seccion(6, f"Señales que persisten ({len(c['persisten'])})") + tabla_senales(c["persisten"], "Ninguna señal persiste.", True)
    cuerpo += seccion(7, "Nota metodológica")
    cuerpo += nota_pdf(c.get("nota") or "El emparejamiento de señales entre documentos es aproximado (similitud de títulos).")
    meta = [("Documento anterior", nom_a), ("Documento actual", nom_b), ("Fecha de emisión", _fecha_hoy()), ("Clasificación", "Confidencial")]
    return pagina_pdf("Comparación de Documentos", "Evolución entre dos versiones analizadas", cuerpo, meta, "Comparación de Documentos")


def excel_plan(d, acciones=None):
    """Excel ordenado con el plan de acción y las señales. Devuelve bytes (.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    acciones = acciones or {}
    wb = Workbook()
    borde = Border(bottom=Side(style="thin", color="C9D0DA"))
    cab_fill = PatternFill("solid", fgColor="1B2F4B")
    zebra = PatternFill("solid", fgColor="F3F5F8")

    def hoja(ws, titulo, encabezados, filas, anchos):
        """Encabezado en la fila 1 (única fila fija), sin títulos arriba: así el scroll funciona normal."""
        for j, h in enumerate(encabezados, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = Font(name="Calibri", bold=True, color="FFFFFF")
            c.fill = cab_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 22
        for i, fila in enumerate(filas):
            maxl = 1
            for j, v in enumerate(fila, 1):
                c = ws.cell(row=2 + i, column=j, value=v)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = borde
                c.font = Font(name="Calibri", size=10)
                if i % 2:
                    c.fill = zebra
                ancho = anchos[j - 1] if j - 1 < len(anchos) else 12
                maxl = max(maxl, -(-len(str(v if v is not None else "")) // max(1, int(ancho * 1.1))))
            ws.row_dimensions[2 + i].height = min(160, 15 * maxl + 3)   # alto fijo: Excel no recalcula al abrir
        for j, a in enumerate(anchos, 1):
            ws.column_dimensions[chr(64 + j)].width = a
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = f"A1:{chr(64 + len(encabezados))}{1 + len(filas)}"
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        ws.oddFooter.center.text = f"{str(d['documento']['nombre']).replace('&', '&&')}  -  Página &P de &N"

    ws = wb.active
    ws.title = "Plan de acción"
    filas = []
    for r in d["recomendaciones"]:
        a = acciones.get(r["id"]) or {}
        est = a.get("estado") if a.get("estado") in _ESTADOS_PLAN else "pendiente"
        filas.append([r["agente_nombre"], r["texto"], a.get("responsable", "") or "", a.get("fecha", "") or "", _ESTADOS_PLAN[est]])
    hoja(ws, "Plan de acción", ["Área", "Acción recomendada", "Responsable", "Fecha límite", "Estado"], filas, [20, 70, 22, 15, 14])

    ws2 = wb.create_sheet("Señales")
    orden_sev = {k: i for i, k in enumerate(SEVERIDADES)}
    sens = sorted(d["senales"], key=lambda x: (orden_sev.get(x["severidad"], 9), x.get("transversal", False)))
    filas = [[x["severidad"].capitalize(), x["agente_nombre"], x["titulo"], x.get("descripcion", "")] for x in sens]
    hoja(ws2, "Señales detectadas", ["Severidad", "Área", "Señal", "Descripción"], filas, [14, 20, 45, 80])

    ws3 = wb.create_sheet("Estado por área")
    filas = []
    for a in sorted(d["agentes"], key=lambda a: (a["score"] is None, a["score"] if a["score"] is not None else 999, a["nombre"])):
        s_ = a["senales"]
        filas.append([a["nombre"], a["score"] if a["con_info"] else None, a["etiqueta"], s_["crítica"], s_["alta"], s_["media"], s_["baja"],
                      a["n_omisiones"], a["n_preguntas"], a["n_recomendaciones"]])
    hoja(ws3, "Estado por área de análisis", ["Área", "Puntaje", "Estado", "Críticas", "Altas", "Medias", "Bajas", "Omisiones", "Preguntas", "Recomend."],
         filas, [20, 10, 26, 10, 10, 10, 10, 11, 11, 11])
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def crear_router(get_supabase):
    router = APIRouter()
    aqui = os.path.dirname(os.path.abspath(__file__))

    @router.get("/api/dashboard/{document_id}")
    def api_dashboard(document_id: str, refrescar: int = 0):
        try:
            data = dashboard_cacheado(get_supabase(), document_id, bool(refrescar))
        except Exception as e:  # nunca romper el servidor por un módulo opcional
            print(f"[DASHBOARD] Error: {e}", flush=True)
            return JSONResponse({"ok": False, "error": str(e)})
        if data is None:
            return JSONResponse({"ok": False, "error": "Documento no encontrado"})
        return data

    @router.get("/api/dashboard-comparar/{id_a}/{id_b}")
    def api_comparar(id_a: str, id_b: str):
        try:
            data = comparar_documentos(get_supabase(), id_a, id_b)
        except Exception as e:
            print(f"[DASHBOARD] Error comparando: {e}", flush=True)
            return JSONResponse({"ok": False, "error": str(e)})
        if data is None:
            return JSONResponse({"ok": False, "error": "Documento no encontrado"})
        return data

    @router.post("/api/dashboard/{document_id}/pdf")
    def api_dashboard_pdf(document_id: str, payload: dict = Body(default={})):
        try:
            data = dashboard_cacheado(get_supabase(), document_id)
            if data is None:
                return JSONResponse({"ok": False, "error": "Documento no encontrado"}, status_code=404)
            pdf = html_a_pdf(html_pdf_dashboard(data, (payload or {}).get("acciones")))
        except Exception as e:
            print(f"[DASHBOARD] Error PDF: {e}", flush=True)
            return JSONResponse({"ok": False, "error": str(e)}, status_code=500)
        nombre = f"Directa_Dashboard_{datetime.now().strftime('%Y-%m-%d')}.pdf"
        return Response(content=pdf, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{nombre}"'})

    @router.post("/api/dashboard-comparar/{id_a}/{id_b}/pdf")
    def api_comparar_pdf(id_a: str, id_b: str):
        try:
            data = comparar_documentos(get_supabase(), id_a, id_b)
            if data is None:
                return JSONResponse({"ok": False, "error": "Documento no encontrado"}, status_code=404)
            pdf = html_a_pdf(html_pdf_comparacion(data))
        except Exception as e:
            print(f"[DASHBOARD] Error PDF comparación: {e}", flush=True)
            return JSONResponse({"ok": False, "error": str(e)}, status_code=500)
        nombre = f"Directa_Comparacion_{datetime.now().strftime('%Y-%m-%d')}.pdf"
        return Response(content=pdf, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{nombre}"'})

    @router.post("/api/dashboard/{document_id}/excel")
    def api_dashboard_excel(document_id: str, payload: dict = Body(default={})):
        try:
            data = dashboard_cacheado(get_supabase(), document_id)
            if data is None:
                return JSONResponse({"ok": False, "error": "Documento no encontrado"}, status_code=404)
            xlsx = excel_plan(data, (payload or {}).get("acciones"))
        except ImportError:
            return JSONResponse({"ok": False, "error": "Falta instalar openpyxl (pip install openpyxl)"}, status_code=500)
        except Exception as e:
            print(f"[DASHBOARD] Error Excel: {e}", flush=True)
            return JSONResponse({"ok": False, "error": str(e)}, status_code=500)
        nombre = f"Directa_Plan_de_Accion_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        return Response(content=xlsx, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{nombre}"'})

    @router.get("/dashboard.html")
    def pagina_dashboard():
        ruta = os.path.join(aqui, "dashboard.html")
        if os.path.exists(ruta):
            return FileResponse(ruta, media_type="text/html")
        return JSONResponse({"ok": False, "error": "dashboard.html no está junto a main.py"}, status_code=404)

    return router