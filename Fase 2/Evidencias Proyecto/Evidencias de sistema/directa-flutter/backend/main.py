import os
import re
import io
import time
import json
import uuid
import shutil
import threading
import traceback
import unicodedata
from functools import lru_cache
from datetime import datetime

from dotenv import load_dotenv

from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse

from supabase import create_client, Client

from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

import docx
import openpyxl

from xhtml2pdf import pisa

from ia_config import (
    generar_embeddings_batch,
    generar_embedding_pregunta,
    llamar_llm_chat,
    llamar_llm_analisis,
    parsear_json_respuesta,
    TAMANO_LOTE_EMBEDDING,
    PAUSA_ENTRE_LOTES,
)

from agentes import obtener_agente, listar_agentes
import agentes as _agentes_mod


load_dotenv()

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL: raise Exception("Falta SUPABASE_URL")
if not SUPABASE_KEY: raise Exception("Falta SUPABASE_KEY")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)


def _es_error_conexion(e):
    t = (type(e).__name__ + " " + str(e)).lower()
    return any(x in t for x in ("disconnected", "remoteprotocol", "connectionterminated", "readerror",
                                "connecterror", "reset by peer", "timed out", "timeout", "closed"))


def _reconectar_supabase():
    """Crea un cliente nuevo. Todo el código usa la variable global `supabase`, así que toma el nuevo solo."""
    global supabase
    try:
        supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
    except Exception as e:
        print(f"[DB] No se pudo reconectar: {e}", flush=True)


def _reintentar_db(fn, intentos=3):
    import functools

    @functools.wraps(fn)
    def envoltura(*a, **k):
        for n in range(1, intentos + 1):
            try:
                return fn(*a, **k)
            except Exception as e:
                if n == intentos or not _es_error_conexion(e):
                    raise
                print(f"[DB] Conexión perdida ({e}); reintentando {n}/{intentos - 1}", flush=True)
                time.sleep(0.6 * n)
                _reconectar_supabase()
    return envoltura

app = FastAPI(title="Directa MVP", version="2.9")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

for _mod in ("dashboard", "memoria", "indicadores", "guia"):
    try:
        _m = __import__(_mod)
        app.include_router(_m.crear_router(lambda: supabase))
    except Exception as _e:
        print(f"[aviso] módulo '{_mod}' no cargado: {_e}")

progreso_documentos = {}
jobs_analisis = {}

CHUNKS_POR_AGENTE = 12
CHUNKS_POR_GRUPO = 30            
CHARS_POR_CHUNK = 1000           

REINTENTAR_ANALISIS = os.getenv("REINTENTAR_ANALISIS", "1").strip().lower() not in ("0", "false", "no")
CONTEXTO_MAX_CHARS = int(os.getenv("CONTEXTO_MAX_CHARS", "9000"))
CONTEXTO_MAX_CHARS_INTEGRADORES = int(os.getenv("CONTEXTO_MAX_CHARS_INTEGRADORES", "6000"))
MAX_TOKENS_GRUPO = int(os.getenv("MAX_TOKENS_GRUPO", "4500"))
MAX_TOKENS_INTEGRADORES = int(os.getenv("MAX_TOKENS_INTEGRADORES", "3500"))
RESUMEN_PREVIO_CHARS = 400       

SCORE_DESDE_SENALES = os.getenv("SCORE_DESDE_SENALES", "1") != "0"
PENALIZACION_SEVERIDAD = {"crítica": 30, "alta": 15, "media": 6, "baja": 2}
SCORE_SIN_SENALES = 90
SCORE_MINIMO = 5

PAUSA_ENTRE_GRUPOS = 3
SEGUNDOS_POR_GRUPO = 60
SEGUNDOS_POR_CHUNK_EMBEDDING = 0.6


GRUPOS_ANALISIS = [
    {"nombre": "Ola 1 · Financiero + Ciberseguridad",
     "agentes": ["financiero", "ciberseguridad"], "recibe_resumenes": False},
    {"nombre": "Ola 2 · Auditoría + Legal + Capital + Operaciones",
     "agentes": ["auditoria", "legal", "capital", "operaciones"], "recibe_resumenes": True},
    {"nombre": "Ola 3 · Comercial + RRHH + Tecnología + Compensación",
     "agentes": ["comercial", "rrhh", "tecnologia", "compensacion"], "recibe_resumenes": True},
    {"nombre": "Ola 4 · ESG + Stakeholders",
     "agentes": ["esg", "stakeholders"], "recibe_resumenes": True},
    {"nombre": "Ola 5 · Integradores: Riesgos + Estrategia",
     "agentes": ["riesgos", "estrategia"], "recibe_resumenes": True,
     "resumen_extendido": True,
     "contexto_chars": CONTEXTO_MAX_CHARS_INTEGRADORES,
     "max_tokens": MAX_TOKENS_INTEGRADORES},
]

TODOS_LOS_AGENTES = [
    "financiero", "ciberseguridad", "auditoria", "legal",
    "capital", "operaciones", "comercial", "rrhh",
    "tecnologia", "compensacion", "esg", "stakeholders",
    "riesgos", "estrategia",
]

KEYWORDS_AGENTE = {
    "financiero": ["ingreso", "ebitda", "margen", "utilidad", "balance", "flujo", "efectivo", "caja", "deuda",
                   "patrimonio", "resultado", "ventas", "costo", "gasto", "presupuesto", "liquidez", "covenant",
                   "apalancamiento", "capital de trabajo", "tipo de cambio", "estado de resultados", "proyeccion",
                   "resultado atribuible", "pérdida del ejercicio", "ganancia bruta", "análisis razonado"],
    "riesgos": ["riesgo", "exposición", "mitigación", "matriz", "contingencia", "amenaza", "impacto",
                "probabilidad", "control", "apetito", "gestión de riesgos", "independiente", "escenario",
                "sensibilidad", "concentración"],
    "ciberseguridad": ["ciberseguridad", "ciber", "seguridad", "incidente", "vulnerabilidad", "pentest", "iso",
                       "nist", "acceso", "firewall", "phishing", "ransomware", "parche", "mfa", "autenticación",
                       "respaldo", "cve", "soc", "controles"],
    "auditoria": ["auditoría", "auditor", "control interno", "hallazgo", "provisión", "material", "revisión",
                  "observación", "deficiencia", "conciliación", "reversión", "carta de control", "salvedad",
                  "recurrente", "estados financieros", "opinión"],
    "legal": ["legal", "litigio", "juicio", "contrato", "regulatorio", "cumplimiento", "multa", "sanción",
              "abogado", "demanda", "contingencia", "probable", "dispensa", "covenant", "tribunal", "laboral",
              "cmf", "superintendencia", "patrimonio mínimo", "filial"],
    "capital": ["capex", "inversión", "roic", "retorno", "proyecto", "tir", "van", "payback", "activo fijo",
                "expansión", "ampliación", "wacc", "costo de capital", "sobrecosto", "desviación", "avance físico"],
    "operaciones": ["operación", "operativo", "planta", "producción", "suministro", "proveedor", "logística",
                    "inventario", "continuidad", "capacidad", "oee", "cobertura", "contrato exclusivo", "spot",
                    "stock", "bodega", "mantención", "paro", "simulacro", "abastecimiento"],
    "comercial": ["venta", "cliente", "pipeline", "comercial", "mercado", "precio", "pricing", "cuota",
                  "segmento", "canal", "competencia", "concentración", "tasa de cierre", "cartera",
                  "participación de mercado", "churn"],
    "rrhh": ["personal", "empleado", "dotación", "rotación", "clima", "talento", "sucesión", "sucesor",
             "capacitación", "remuneración", "sindicato", "rrhh", "recursos humanos", "ausentismo",
             "negociación colectiva", "cargo clave", "colaboradores", "huelga"],
    "tecnologia": ["tecnología", "sistema", "software", "infraestructura", "ti", "digital", "proyecto ti",
                   "obsolescencia", "proveedor ti", "plataforma", "erp", "crm", "servidor", "migración",
                   "soporte del fabricante", "nube", "licencia", "implementación", "go-live"],
    "compensacion": ["compensación", "incentivo", "bono", "sueldo", "benchmark", "meta ejecutiva",
                     "remuneración variable", "percentil", "política de compensación", "clawback",
                     "plan de incentivos", "remuneración", "gerencia"],
    "esg": ["esg", "sostenibilidad", "emisión", "carbono", "ambiental", "climático", "ncg 461", "social",
            "gobernanza", "huella", "alcance 1", "alcance 2", "alcance 3", "estrés hídrico", "transición",
            "descarbonización", "tcfd", "diversidad", "residuos", "energía"],
    "stakeholders": ["stakeholder", "grupos de interés", "reputación", "comunidad", "regulador", "prensa",
                     "accionista", "relacionamiento", "medio", "huelga", "sindicato", "carta", "conflicto",
                     "vecinos", "mitigación", "negociación colectiva"],
    "estrategia": ["estrategia", "estratégico", "plan", "visión", "misión", "objetivo", "meta", "kpi", "rumbo",
                   "transformación", "pilar", "avance", "hoja de ruta", "iniciativa", "horizonte"],
    "general": [],
}


def guardar_documento(document_id, nombre_archivo, total_paginas):
    existe = supabase.table("documentos").select("id").eq("id", document_id).execute()
    if len(existe.data) == 0:
        supabase.table("documentos").insert({
            "id": document_id,
            "titulo": nombre_archivo,
            "nombre_archivo": nombre_archivo,
            "total_paginas": total_paginas,
            "estado": "Procesando",
        }).execute()


def actualizar_estado(document_id, estado):
    try:
        supabase.table("documentos").update({"estado": estado}).eq("id", document_id).execute()
    except Exception as e:
        print(f"[ESTADO] No se pudo actualizar: {e}", flush=True)


def _contar_chunks_doc(document_id):
    try:
        r = supabase.table("document_chunks").select("id", count="exact").eq("document_id", document_id).limit(1).execute()
        return r.count or 0
    except Exception:
        return None


def guardar_chunks_batch(document_id, items):
    if not items:
        return
    payload = [
        {"document_id": document_id, "content": it["content"], "page_number": it["page_number"], "entity": None, "embedding": it["embedding"]}
        for it in items
    ]
    n0 = _contar_chunks_doc(document_id)
    for intento in (1, 2, 3):
        try:
            supabase.table("document_chunks").insert(payload).execute()
            return
        except Exception as e:
            if intento == 3 or not _es_error_conexion(e):
                raise
            print(f"[DB] Conexión perdida guardando lote ({e}); reintentando {intento}/2", flush=True)
            time.sleep(0.8 * intento)
            _reconectar_supabase()
            n1 = _contar_chunks_doc(document_id)
            if n0 is not None and n1 is not None and n1 >= n0 + len(payload):
                return  # el lote sí se había guardado: no se repite


def _obtener_analisis_guardado(document_id, agente_key):
    try:
        resp = (supabase.table("analisis").select("resultado")
                .eq("document_id", document_id).eq("agente", agente_key).limit(1).execute())
        if resp.data:
            return resp.data[0]["resultado"]
    except Exception as e:
        print(f"[CACHE] Error leyendo {agente_key}: {e}", flush=True)
    return None


def _guardar_analisis(document_id, agente_key, resultado):
    for intento in (1, 2, 3):
        try:
            supabase.table("analisis").delete().eq("document_id", document_id).eq("agente", agente_key).execute()
            supabase.table("analisis").insert({
                "document_id": document_id, "agente": agente_key,
                "resultado": resultado, "score": resultado.get("score", 0),
            }).execute()
            return
        except Exception as e:
            if intento < 3 and _es_error_conexion(e):
                print(f"[DB] Conexión perdida guardando {agente_key}; reintentando", flush=True)
                time.sleep(0.8 * intento)
                _reconectar_supabase()
                continue
            print(f"[CACHE] No se pudo guardar {agente_key}: {e}", flush=True)
            return


def _obtener_todos_los_analisis(document_id):
    for intento in (1, 2, 3):
        try:
            resp = (supabase.table("analisis").select("agente, resultado")
                    .eq("document_id", document_id).execute())
            return {row["agente"]: row["resultado"] for row in (resp.data or [])}
        except Exception as e:
            if intento < 3 and _es_error_conexion(e):
                print(f"[DB] Conexión perdida leyendo análisis; reintentando", flush=True)
                time.sleep(0.8 * intento)
                _reconectar_supabase()
                continue
            print(f"[CACHE] Error leyendo análisis: {e}", flush=True)
            return {}


def _normalizar_layout(texto):
    """El modo 'layout' de pypdf mantiene las tablas alineadas en columnas con muchos espacios.
    Se convierten esas separaciones en ' | ' para que cada fila quede como
    'Acero galvanizado | 78% | 1 contrato exclusivo + spot | 22 | 30' y el modelo no confunda
    qué número va con qué encabezado (el modo normal deja cada celda en una línea suelta)."""
    lineas = []
    for linea in (texto or "").splitlines():
        linea = re.sub(r"[ \t]{3,}", " | ", linea.strip())
        linea = re.sub(r"[ \t]{2,}", " ", linea)
        lineas.append(linea)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lineas)).strip()


_CELDA_NUM = re.compile(r"^[(\-\u2013+]?\$?\d[\d.,]*\)?%?$|^[-\u2013\u2014]$")


def _filas_tabla(texto_layout):
    """Del texto en modo 'layout' toma SOLO las filas que parecen de una tabla (una etiqueta y
    2 o más cifras en columnas) y las deja como 'Etiqueta | 2,5 | 8,1 | (5,5)', con la fila de
    encabezados que las precede. La prosa NO se toma de aquí: en el modo layout los textos
    justificados salen con letras separadas y las memorias de dos o tres columnas se mezclan
    renglón a renglón (para la prosa se usa el modo normal)."""
    filas, previa, en_bloque = [], None, False
    for linea in (texto_layout or "").splitlines():
        partes = [x.strip() for x in re.split(r"[ \t]{3,}", linea.strip()) if x.strip()]
        nums = sum(1 for x in partes[1:] if _CELDA_NUM.match(x))
        es_fila = (3 <= len(partes) <= 12 and not _CELDA_NUM.match(partes[0])
                   and re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{2}", partes[0]) is not None
                   and nums >= 2 and nums * 2 >= len(partes) - 1
                   and all(len(x) <= 40 for x in partes[1:] if not _CELDA_NUM.match(x)))
        if es_fila:
            if not en_bloque and previa and len(previa) >= 2:
                filas.append(" | ".join(previa))
            filas.append(" | ".join(partes))
            en_bloque = True
        else:
            en_bloque = False
        previa = partes if len(partes) >= 2 else None
    return "\n".join(filas)


def _quitar_encabezados_repetidos(paginas):
    """Quita los encabezados y pies que se repiten en muchas páginas ('Compañía X S.A.', 'Análisis
    Razonado', 'Memoria Anual 2017'…). Gastaban un 5-10% del texto que ve cada agente. Solo se
    eliminan líneas de 12+ letras que aparecen en al menos el 20% de las páginas (mínimo 8)."""
    if len(paginas) < 12:
        return paginas

    def clave(linea):
        return re.sub(r"\d+", "#", " ".join(linea.split())).strip().lower()

    conteo = {}
    for texto, _ in paginas:
        for k in {clave(l) for l in texto.splitlines() if 12 <= len(l.strip()) <= 120 and len(re.findall(r"\d+", l)) <= 3}:
            conteo[k] = conteo.get(k, 0) + 1
    minimo = max(8, int(len(paginas) * 0.20))
    repetidas = {k for k, n in conteo.items()
                 if n >= minimo and sum(c.isalpha() for c in k) >= 12 and not k.startswith("[tabla")}
    if not repetidas:
        return paginas
    limpias = []
    for texto, nro in paginas:
        lineas = [l for l in texto.splitlines()
                  if not (12 <= len(l.strip()) <= 120 and len(re.findall(r"\d+", l)) <= 3 and clave(l) in repetidas)]
        limpias.append((re.sub(r"\n{3,}", "\n\n", "\n".join(lineas)).strip(), nro))
    return limpias


def _cargar_pdf(ruta):
    """Carga el PDF. El texto de la página sale del modo normal (respeta las columnas y no
    rompe las palabras). Si la página trae tablas, se agrega debajo un bloque con sus filas
    alineadas con ' | ' (modo layout) para que el modelo no confunda qué cifra va con qué
    encabezado. Si el modo layout no está disponible o falla, queda solo el texto normal."""
    paginas = PyPDFLoader(ruta).load()
    base = [(p.page_content or "", int(p.metadata.get("page", 0)) + 1) for p in paginas]
    try:
        lay = PyPDFLoader(ruta, extraction_mode="layout").load()
    except TypeError:
        print("[PDF] extraction_mode='layout' no disponible; solo texto normal", flush=True)
        return base
    except Exception as e:
        print(f"[PDF] Modo layout falló ({e}); solo texto normal", flush=True)
        return base
    por_pagina = {}
    for p in lay:
        try:
            tabla = _filas_tabla(p.page_content)
        except Exception:
            tabla = ""
        if tabla:
            por_pagina[int(p.metadata.get("page", 0)) + 1] = tabla
    salida = []
    for texto, nro in base:
        if nro in por_pagina:
            texto = texto.rstrip() + "\n\n[Tabla de esta página, columnas separadas con |]\n" + por_pagina[nro]
        salida.append((texto, nro))
    return salida


def _extraer_texto_archivo(ruta, nombre_archivo):
    ext = nombre_archivo.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        paginas = _quitar_encabezados_repetidos(_cargar_pdf(ruta))
        return paginas, len(paginas)

    elif ext in ("docx", "doc"):
        documento = docx.Document(ruta)
        partes = []
        for p in documento.paragraphs:
            if p.text.strip():
                partes.append(p.text.strip())
        for tabla in documento.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells]
                if any(celdas):
                    partes.append(" | ".join(celdas))
        texto_completo = "\n\n".join(partes)
        bloques = []
        for i in range(0, len(texto_completo), 3000):
            bloque = texto_completo[i:i+3000]
            pagina = i // 3000 + 1
            bloques.append((bloque, pagina))
        return bloques, len(bloques)

    elif ext in ("xlsx", "xls"):
        wb = openpyxl.load_workbook(ruta, data_only=True)
        bloques = []
        for idx, hoja in enumerate(wb.worksheets, start=1):
            lineas = []
            for row in hoja.iter_rows(values_only=True):
                valores = [str(c) if c is not None else "" for c in row]
                if any(valores):
                    lineas.append(" | ".join(valores))
            texto_hoja = f"[Hoja: {hoja.title}]\n" + "\n".join(lineas)
            bloques.append((texto_hoja, idx))
        return bloques, len(bloques)

    else:
        raise Exception(f"Formato no soportado: {ext}")


def _tiene_contenido_sustantivo(texto):
    if len(texto) < 50:
        return False
    digitos = sum(1 for c in texto if c.isdigit())
    letras = sum(1 for c in texto if c.isalpha())
    return digitos >= 5 or letras >= 100


def _normalizar_texto(texto):
    """Minúsculas y sin tildes, para comparar palabras clave de forma robusta."""
    t = unicodedata.normalize("NFD", (texto or "").lower())
    return "".join(ch for ch in t if unicodedata.category(ch) != "Mn")


@lru_cache(maxsize=None)
def _patron_keyword(kw):
    """La palabra clave debe empezar en límite de palabra; las cortas (ti, tir, van, iso, erp...)
    también deben terminar en límite, para no coincidir dentro de 'estimar', 'activo', etc."""
    k = re.escape(_normalizar_texto(kw)).replace("\\ ", r"\s+")
    fin = r"(?!\w)" if len(kw) <= 4 else ""
    return re.compile(r"(?<!\w)" + k + fin)


def _hits_keywords(texto_norm, keywords):
    """3 puntos por palabra clave presente (+1 por repetición, hasta +2)."""
    puntos = 0
    for kw in keywords:
        n = len(_patron_keyword(kw).findall(texto_norm))
        if n:
            puntos += 3 + min(2, n - 1)
    return puntos


def _bonus_chunk(texto):
    digitos = sum(1 for c in texto if c.isdigit())
    return min(5, digitos // 20) + (2 if _tiene_contenido_sustantivo(texto) else 0)


def _puntuar_chunk(chunk, keywords, texto_norm=None):
    texto = texto_norm if texto_norm is not None else _normalizar_texto(chunk.get("content") or "")
    return _hits_keywords(texto, keywords) + _bonus_chunk(texto)


def _cargar_chunks_documento(document_id, tam_pagina=1000, max_total=30000):
    """Lee TODOS los fragmentos del documento, en orden de lectura (página, id).
    Supabase entrega como máximo 1000 filas por consulta, por eso se pagina."""
    filas_totales = []
    desde = 0
    while desde < max_total:
        resp = (supabase.table("document_chunks").select("id, content, page_number")
                .eq("document_id", document_id)
                .order("page_number").order("id")
                .range(desde, desde + tam_pagina - 1).execute())
        filas = resp.data or []
        filas_totales.extend(filas)
        if len(filas) < tam_pagina:
            break
        desde += tam_pagina
    return filas_totales


def _chunks_utiles(chunks):
    utiles = [c for c in chunks if _tiene_contenido_sustantivo(c.get("content") or "")]
    return utiles if len(utiles) >= 2 else chunks


def _preparar_chunks(chunks):
    todos = _chunks_utiles(chunks)
    textos = [_normalizar_texto(c.get("content") or "") for c in todos]
    bonus = [_bonus_chunk(t) for t in textos]
    return todos, textos, bonus


def _ranking_agente(textos, bonus, agente_key):
    """Índices de fragmentos ordenados de mejor a peor para el agente.
    Solo entran los que mencionan al menos una de sus palabras clave."""
    keywords = KEYWORDS_AGENTE.get(agente_key, [])
    puntuados = []
    for i, t in enumerate(textos):
        hits = _hits_keywords(t, keywords)
        if hits > 0:
            puntuados.append((-(hits + bonus[i]), i))
    puntuados.sort()
    return [i for _, i in puntuados]


def seleccionar_chunks_para_agente(document_id, agente_key, max_chunks=CHUNKS_POR_AGENTE, _chunks=None):
    chunks = _chunks if _chunks is not None else _cargar_chunks_documento(document_id)
    todos, textos, bonus = _preparar_chunks(chunks)
    if not todos:
        return []
    ranking = _ranking_agente(textos, bonus, agente_key)[:max_chunks]
    return [todos[i] for i in sorted(ranking)]


def _costo_chunk(c):
    return min(len(c.get("content") or ""), CHARS_POR_CHUNK) + 16


def seleccionar_chunks_para_grupo(document_id, agentes_keys, max_chunks=CHUNKS_POR_GRUPO, max_chars=None):
    """Elige los fragmentos de la ola repartiendo el presupuesto de caracteres entre los
    agentes (uno por turno: el mejor de cada uno, luego el segundo mejor, etc.).
    Antes se tomaban los primeros 72 fragmentos y luego se recortaba por número de página,
    lo que dejaba fuera las páginas finales del informe."""
    presupuesto = int(max_chars or CONTEXTO_MAX_CHARS)
    todos, textos, bonus = _preparar_chunks(_cargar_chunks_documento(document_id))
    if not todos:
        return []
    rankings = [_ranking_agente(textos, bonus, k) for k in agentes_keys]
    posiciones = [0] * len(rankings)
    elegidos = set()
    usados = 0
    avanzo = True
    while avanzo and len(elegidos) < max_chunks:
        avanzo = False
        for n, ranking in enumerate(rankings):
            while posiciones[n] < len(ranking):
                i = ranking[posiciones[n]]
                posiciones[n] += 1
                if i in elegidos:
                    continue
                costo = _costo_chunk(todos[i])
                if usados + costo > presupuesto:
                    continue
                elegidos.add(i)
                usados += costo
                avanzo = True
                break
            if len(elegidos) >= max_chunks:
                break
    return [todos[i] for i in sorted(elegidos)]


def _reglas_base_analisis():
    """Reglas de verificación comunes (definidas en agentes.py). Se agregan UNA sola vez por
    prompt de ola. Si agentes.py ya las incrusta en cada rol, no se repiten."""
    ya_incluidas = getattr(_agentes_mod, "INCLUIR_REGLAS_BASE_ANALISIS",
                           getattr(_agentes_mod, "INCLUIR_REGLAS_BASE", False))
    if ya_incluidas:
        return ""
    return (getattr(_agentes_mod, "REGLAS_BASE_ANALISIS", "") or "").strip()


def _construir_prompt_grupo(agentes_keys, resumenes_previos=None, integrador=False):
    partes = []
    partes.append(
        "Eres un PANEL de analistas senior preparando a un Gerente General "
        "para una sesión de directorio. Debes responder como CADA UNO de los "
        "siguientes agentes especializados, manteniendo estrictamente su voz, "
        "su foco y sus criterios particulares."
    )
    partes.append("\n\nAGENTES EN ESTE PANEL:\n")
    for key in agentes_keys:
        config = obtener_agente(key)
        partes.append(f"\n### {key.upper()} — {config['nombre']}\n{config['rol_analisis']}\n")

    partes.append("""
REGLAS CRÍTICAS:

1. Cada agente responde SOLO desde su especialidad. NO mezcles voces.

2. PASO 1 — EVALÚA LA RELEVANCIA:
   Antes de generar hallazgos, pregúntate: "¿El documento contiene información
   sustantiva sobre MI especialidad?"

   - Si el documento es una tabla de sueldos, una carta corta o similar, NO hay info.
   - Si la respuesta es NO:
     * "resumen_ejecutivo": "El documento no contiene información sustantiva
       para el área de [tu especialidad]. Su contenido corresponde a [describe]."
     * "hallazgos": []
     * "senales": []
     * "omisiones": []
     * "preguntas": []
     * "recomendaciones": []
     * "score": 90
     NO INVENTES problemas.

3. PASO 2 — SI HAY INFORMACIÓN RELEVANTE:
   - resumen_ejecutivo: 3-5 líneas, con las cifras clave de tu área
   - hallazgos: los hechos relevantes que el documento respalda (normalmente 2 a 5)
   - senales: SOLO situaciones que el documento respalda y que superan el umbral de tu
     especialidad. Si no hay ninguna, deja la lista vacía. NO rellenes para completar.
   - omisiones: SOLO información que realmente no aparece en los fragmentos entregados.
     Si el dato está en una tabla o nota, no es una omisión. Lista vacía si no hay.
   - preguntas: 2 a 4, concretas, indicando el dato o hecho que las motiva
   - recomendaciones: una por cada señal u omisión relevante (lista vacía si no hay)
   - score: 0-100 (orientativo)

4. NO uses markdown. Texto plano.

5. Basa tus hallazgos en CONTENIDO REAL. NO inventes cifras, fechas ni números de
   página. Cuando cites una página, usa solo las que aparecen como [Página N] en los
   fragmentos.
""")

    partes.append(
        f"\nFECHA DE HOY: {datetime.now().strftime('%d/%m/%Y')}. Toma como fecha de corte la que indique el "
        "propio documento. Cualquier plazo que recomiendes debe ser POSTERIOR a esa fecha de corte y a hoy "
        "(nunca propongas un plazo que ya venció)."
    )
    reglas_base = _reglas_base_analisis()
    if reglas_base:
        partes.append("\nREGLAS DE VERIFICACIÓN COMUNES A TODOS LOS AGENTES:\n" + reglas_base + "\n")

    if resumenes_previos:
        if integrador:
            partes.append(
                "\n\nCONTEXTO: RESÚMENES Y SEÑALES PRINCIPALES DE LOS AGENTES YA ANALIZADOS "
                "(tu trabajo es cruzarlos entre sí y con los fragmentos; cita siempre qué agente "
                "o qué hecho respalda cada correlación y no inventes hallazgos que no estén aquí)\n"
            )
        else:
            partes.append("\n\nCONTEXTO: RESÚMENES DE AGENTES QUE YA ANALIZARON\n")
        for k, resumen in resumenes_previos.items():
            partes.append(f"\n[{k.upper()}] {resumen}\n")

    partes.append("\n\nRESPONDE ÚNICAMENTE CON ESTE JSON (una clave por agente):\n")
    ejemplo = {}
    for key in agentes_keys:
        ejemplo[key] = {
            "resumen_ejecutivo": "string",
            "hallazgos": [{"titulo": "string", "descripcion": "string"}],
            "senales": [{"titulo": "string", "descripcion": "string", "severidad": "crítica|alta|media|baja"}],
            "omisiones": [{"descripcion": "string"}],
            "preguntas": [{"pregunta": "string"}],
            "recomendaciones": [{"descripcion": "string"}],
            "score": 0,
        }
    partes.append(json.dumps(ejemplo, indent=2, ensure_ascii=False))
    return "\n".join(partes)


def _construir_contexto_analisis(chunks):
    return "\n\n".join(
        f"[Página {c.get('page_number') or '?'}]\n{(c.get('content') or '')[:CHARS_POR_CHUNK]}"
        for c in chunks
    )


_ORDEN_SEV = {"crítica": 0, "alta": 1, "media": 2, "baja": 3}


def _norm_sev(valor):
    """Devuelve siempre 'crítica' | 'alta' | 'media' | 'baja'. Si el modelo copió la plantilla
    ('crítica|alta|media|baja') o escribió algo irreconocible, se trata como 'media'."""
    t = _normalizar_texto(str(valor or ""))
    if "|" in t:
        return "media"
    if "crit" in t:
        return "crítica"
    if "alt" in t:
        return "alta"
    if "baj" in t:
        return "baja"
    return "media"


def _calcular_score(resultado):
    """Score determinista: 100 menos una penalización por cada señal según su severidad.
    Sin señales = 90 (sin observaciones). Así el número siempre es coherente con lo que se muestra."""
    senales = [s for s in (resultado.get("senales") or []) if isinstance(s, dict)]
    if not senales:
        return SCORE_SIN_SENALES
    penal = sum(PENALIZACION_SEVERIDAD[_norm_sev(s.get("severidad"))] for s in senales)
    return max(SCORE_MINIMO, 100 - penal)


def _analisis_valido(r):
    """True si el análisis tiene contenido real (no es un dict vacío por una respuesta fallida)."""
    if not isinstance(r, dict):
        return False
    if str(r.get("resumen_ejecutivo") or "").strip():
        return True
    return any(r.get(k) for k in ("hallazgos", "senales", "omisiones", "preguntas", "recomendaciones"))


def _resumen_extendido(r):
    """Resumen corto + las 3 señales más graves, para que los integradores puedan correlacionar."""
    partes = [str(r.get("resumen_ejecutivo") or "").strip()[:300]]
    senales = [s for s in (r.get("senales") or []) if isinstance(s, dict) and s.get("titulo")]
    senales.sort(key=lambda s: _ORDEN_SEV[_norm_sev(s.get("severidad"))])
    top = [f"[{_norm_sev(s.get('severidad'))}] {str(s.get('titulo')).strip()[:90]}" for s in senales[:3]]
    if top:
        partes.append("Señales: " + "; ".join(top))
    return " ".join(p for p in partes if p)


def _normalizar_lista(lista, campo_principal):
    if not isinstance(lista, list):
        return []
    out = []
    for item in lista:
        if isinstance(item, dict):
            out.append(item)
        elif isinstance(item, str):
            out.append({campo_principal: item.strip()})
        else:
            out.append({campo_principal: str(item)})
    return out


def _limpiar_unicode(texto):
    """Solo reemplaza caracteres Unicode problemáticos. NO toca markdown."""
    if not isinstance(texto, str):
        return str(texto)
    reemplazos = {
        "\u2011": "-", "\u2013": "-", "\u2014": "-", "\u2015": "-",
        "\u2022": "*", "\u25aa": "*", "\u25a0": "*", "\u25cf": "*", "\u25e6": "*",
        "\u2192": "->", "\u2190": "<-",
        "\u2264": "<=", "\u2265": ">=", "\u2260": "!=",
        "\u00a0": " ",
        "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"',
        "\u2026": "...",
        "\u2122": "(TM)", "\u00ae": "(R)", "\u00a9": "(C)",
    }
    for k, v in reemplazos.items():
        texto = texto.replace(k, v)
    return texto.strip()


def _limpiar_markdown(texto):
    """Limpieza completa (Unicode + markdown). Para ALMACENAMIENTO."""
    if not isinstance(texto, str):
        return str(texto)
    texto = _limpiar_unicode(texto)
    texto = texto.replace("**", "").replace("__", "")
    texto = texto.replace(" -- ", " - ")
    texto = re.sub(r'^\s*#+\s*', '', texto, flags=re.M)
    return texto.strip()


def _esc_html(s):
    """Escapa SOLO < > y &. Las tildes se convierten al final del pipeline."""
    if not isinstance(s, str):
        return str(s)
    s = s.replace("\u25a0", "").replace("\u3010", "(").replace("\u3011", ")")
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _esc_html_ligero(s):
    """Escapa < > & pero permite <b>, <i>, <br/>. NO convierte tildes."""
    if not isinstance(s, str):
        return str(s)
    s = s.replace("<b>", "\x00B\x00").replace("</b>", "\x00/B\x00")
    s = s.replace("<i>", "\x00I\x00").replace("</i>", "\x00/I\x00")
    s = s.replace("<br>", "\x00BR\x00").replace("<br/>", "\x00BR\x00")
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = s.replace("\x00B\x00", "<b>").replace("\x00/B\x00", "</b>")
    s = s.replace("\x00I\x00", "<i>").replace("\x00/I\x00", "</i>")
    s = s.replace("\x00BR\x00", "<br/>")
    return s


def _entidades_para_pdf(html):
    """Barrido FINAL del HTML antes de generar el PDF.
    Convierte tildes a entidades y elimina caracteres raros.
    Se llama UNA SOLA VEZ al final."""
    if not isinstance(html, str):
        return html

    html = html.replace("\u2011", "-").replace("\u2010", "-")
    html = html.replace("\u2012", "-").replace("\u2013", "-")
    html = html.replace("\u2014", "-").replace("\u2015", "-")
    html = html.replace("\u2022", "&bull;").replace("\u2023", "&bull;")
    html = html.replace("\u25aa", "&bull;").replace("\u25a0", "&bull;")
    html = html.replace("\u25cf", "&bull;").replace("\u25e6", "&bull;")
    html = html.replace("\u2192", "-&gt;").replace("\u2190", "&lt;-")
    html = html.replace("\u2264", "&lt;=").replace("\u2265", "&gt;=").replace("\u2260", "!=")
    html = html.replace("\u00a0", " ").replace("\u2026", "...")
    html = html.replace("\u2122", "(TM)").replace("\u00ae", "(R)").replace("\u00a9", "(C)")
    html = html.replace("\u2018", "'").replace("\u2019", "'")
    html = html.replace("\u201c", '"').replace("\u201d", '"')

    reemplazos_acentos = {
        "á": "&aacute;", "é": "&eacute;", "í": "&iacute;", "ó": "&oacute;", "ú": "&uacute;",
        "Á": "&Aacute;", "É": "&Eacute;", "Í": "&Iacute;", "Ó": "&Oacute;", "Ú": "&Uacute;",
        "à": "&agrave;", "è": "&egrave;", "ì": "&igrave;", "ò": "&ograve;", "ù": "&ugrave;",
        "À": "&Agrave;", "È": "&Egrave;", "Ì": "&Igrave;", "Ò": "&Ograve;", "Ù": "&Ugrave;",
        "ä": "&auml;", "ë": "&euml;", "ï": "&iuml;", "ö": "&ouml;", "ü": "&uuml;",
        "Ä": "&Auml;", "Ë": "&Euml;", "Ï": "&Iuml;", "Ö": "&Ouml;", "Ü": "&Uuml;",
        "ñ": "&ntilde;", "Ñ": "&Ntilde;",
        "ç": "&ccedil;", "Ç": "&Ccedil;",
        "¿": "&iquest;", "¡": "&iexcl;",
        "º": "&ordm;", "ª": "&ordf;",
    }
    for k, v in reemplazos_acentos.items():
        html = html.replace(k, v)

    html = re.sub(r'[\x80-\xff]', '', html)
    return html


def _normalizar_resultado_analisis(resultado):
    if not isinstance(resultado, dict):
        resultado = {}
    resultado["hallazgos"] = _normalizar_lista(resultado.get("hallazgos"), "titulo")
    resultado["senales"] = _normalizar_lista(resultado.get("senales"), "titulo")
    resultado["omisiones"] = _normalizar_lista(resultado.get("omisiones"), "descripcion")
    resultado["preguntas"] = _normalizar_lista(resultado.get("preguntas"), "pregunta")
    resultado["recomendaciones"] = _normalizar_lista(resultado.get("recomendaciones"), "descripcion")
    resultado["resumen_ejecutivo"] = _limpiar_markdown(resultado.get("resumen_ejecutivo", ""))
    for lista in ["hallazgos", "senales", "omisiones", "preguntas", "recomendaciones"]:
        for item in resultado[lista]:
            for k, v in list(item.items()):
                if isinstance(v, str):
                    item[k] = _limpiar_markdown(v)
    for s in resultado["senales"]:
        s["severidad"] = _norm_sev(s.get("severidad"))
    score = resultado.get("score", 70)
    try:
        score = int(score)
    except (ValueError, TypeError):
        score = 70
    score = max(0, min(100, score))
    if SCORE_DESDE_SENALES:
        resultado["score_modelo"] = score
        score = _calcular_score(resultado)
    resultado["score"] = score
    return resultado

def _partir_fila_markdown(linea):
    linea = linea.strip()
    if linea.startswith("|"): linea = linea[1:]
    if linea.endswith("|"): linea = linea[:-1]
    return [c.strip() for c in linea.split("|")]


def _celda_md_a_html(texto):
    """Convierte **bold** e *italic* dentro de una celda de tabla a HTML."""
    if not isinstance(texto, str):
        return str(texto)
    texto = texto.replace("<br>", " ").replace("<br/>", " ").replace("\n", " ")
    texto = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', texto)
    texto = re.sub(r'(?<!\*)\*([^*\n]+?)\*(?!\*)', r'<i>\1</i>', texto)
    return _esc_html_ligero(texto)


def _markdown_tablas_a_html(texto):
    """Convierte tablas markdown a HTML con atributos clásicos para xhtml2pdf."""
    if not isinstance(texto, str) or "|" not in texto:
        return texto

    lineas = texto.split("\n")
    salida = []
    i = 0

    while i < len(lineas):
        linea = lineas[i].strip()

        es_tabla = (
            "|" in linea
            and i + 1 < len(lineas)
            and re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', lineas[i + 1])
            and "-" in lineas[i + 1]
        )

        if es_tabla:
            filas = []
            cabecera = _partir_fila_markdown(linea)
            j = i + 2
            while j < len(lineas) and "|" in lineas[j]:
                fila = _partir_fila_markdown(lineas[j])
                if fila:
                    filas.append(fila)
                j += 1

            if cabecera:
                html = '<table border="0" cellpadding="3" cellspacing="0" style="width:100%;margin:4px 0 8px 0;font-size:8.5pt;">'
                html += '<thead><tr bgcolor="#1B2F4B">'
                for col in cabecera:
                    html += f'<th style="color:#fff;text-align:left;padding:3px 5px;font-size:8pt;">{_celda_md_a_html(col)}</th>'
                html += '</tr></thead><tbody>'
                for idx, fila in enumerate(filas):
                    bg = "#ffffff" if idx % 2 == 0 else "#F3F5F8"
                    html += f'<tr bgcolor="{bg}">'
                    for celda in fila:
                        html += f'<td style="padding:3px 5px;color:#2B2F36;border-bottom:0.5pt solid #C9D0DA;vertical-align:top;">{_celda_md_a_html(celda)}</td>'
                    html += '</tr>'
                html += '</tbody></table>'
                salida.append(html)
                i = j
                continue

        salida.append(linea)
        i += 1

    return "\n".join(salida)


def _render_markdown_a_html(texto):
    """Convierte markdown a HTML. NO usa _limpiar_markdown para no perder
    los headers y bold markers antes de procesarlos."""
    if not isinstance(texto, str):
        return str(texto)

    texto = _limpiar_unicode(texto)

    if re.search(r'\|\s*[\-:]+\s*\|', texto):
        texto = _markdown_tablas_a_html(texto)

    lineas = texto.split("\n")
    salida = []
    en_lista_ul = False
    en_lista_ol = False

    def cerrar_listas():
        nonlocal en_lista_ul, en_lista_ol
        if en_lista_ul:
            salida.append("</ul>")
            en_lista_ul = False
        if en_lista_ol:
            salida.append("</ol>")
            en_lista_ol = False

    for linea in lineas:
        l = linea.rstrip()
        ls = l.strip()

        if any(ls.startswith(t) for t in ["<table", "</table", "<thead", "</thead", "<tbody", "</tbody", "<tr", "</tr", "<th", "</th", "<td", "</td"]):
            cerrar_listas()
            salida.append(l)
            continue

        m = re.match(r'^(\s*)(#{1,6})\s+(.+)$', l)
        if m:
            cerrar_listas()
            nivel = len(m.group(2))
            size = {1: 14, 2: 13, 3: 12, 4: 11, 5: 10, 6: 10}.get(nivel, 11)
            titulo = re.sub(r'\*\*(.+?)\*\*', r'\1', m.group(3))
            salida.append(
                f'<div style="font-size:{size}pt;font-weight:700;color:#1B2F4B;'
                f'margin:14px 0 6px;letter-spacing:0.3px;">{_esc_html_ligero(titulo)}</div>'
            )
            continue

        if re.match(r'^[-*_]{3,}$', ls):
            cerrar_listas()
            salida.append('<hr style="border:none;border-top:1px solid #ddd;margin:12px 0;"/>')
            continue

        if ls.endswith(":") and len(ls) > 5 and ls.upper() == ls and any(c.isalpha() for c in ls):
            cerrar_listas()
            salida.append(
                f'<div style="font-size:10pt;font-weight:700;color:#1B2F4B;'
                f'letter-spacing:1px;margin:12px 0 6px;">{_esc_html_ligero(ls)}</div>'
            )
            continue

        m = re.match(r'^\s*[\-\*]\s+(.+)$', l)
        if m:
            if en_lista_ol:
                cerrar_listas()
            if not en_lista_ul:
                salida.append('<ul style="margin:6px 0 10px 22px;padding:0;">')
                en_lista_ul = True
            contenido = m.group(1)
            contenido = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', contenido)
            contenido = re.sub(r'(?<!\*)\*([^*\n]+?)\*(?!\*)', r'<i>\1</i>', contenido)
            salida.append(f'<li style="margin-bottom:5px;color:#222;line-height:1.55;">{_esc_html_ligero(contenido)}</li>')
            continue

        m = re.match(r'^\s*(\d+)\.\s+(.+)$', l)
        if m:
            if en_lista_ul:
                cerrar_listas()
            if not en_lista_ol:
                salida.append('<ol style="margin:6px 0 10px 22px;padding:0;">')
                en_lista_ol = True
            contenido = m.group(2)
            contenido = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', contenido)
            contenido = re.sub(r'(?<!\*)\*([^*\n]+?)\*(?!\*)', r'<i>\1</i>', contenido)
            salida.append(f'<li style="margin-bottom:5px;color:#222;line-height:1.55;">{_esc_html_ligero(contenido)}</li>')
            continue

        if not ls:
            cerrar_listas()
            salida.append('<div style="height:6px;"></div>')
            continue

        cerrar_listas()
        l2 = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', l)
        l2 = re.sub(r'(?<!\*)\*([^*\n]+?)\*(?!\*)', r'<i>\1</i>', l2)
        salida.append(f'<p style="margin:6px 0;color:#222;line-height:1.55;">{_esc_html_ligero(l2)}</p>')

    cerrar_listas()
    return "\n".join(salida)

def _resumenes_previos_para_grupo(document_id, idx_grupo, memoria=None):
    """Resúmenes de los agentes de las olas ANTERIORES (desde la base de datos, más lo recién
    calculado en memoria). Las olas integradoras reciben además sus señales principales."""
    grupo = GRUPOS_ANALISIS[idx_grupo]
    if not grupo.get("recibe_resumenes"):
        return None
    previos = [k for g in GRUPOS_ANALISIS[:idx_grupo] for k in g["agentes"]]
    if not previos:
        return None
    guardados = _obtener_todos_los_analisis(document_id)
    if memoria:
        guardados = {**guardados, **memoria}
    extendido = bool(grupo.get("resumen_extendido"))
    out = {}
    for k in previos:
        r = guardados.get(k)
        if not _analisis_valido(r):
            continue
        if extendido:
            txt = _resumen_extendido(r)
        else:
            txt = str(r.get("resumen_ejecutivo") or "").strip()[:RESUMEN_PREVIO_CHARS]
        if txt:
            out[k] = txt
    return out or None


def _dividir_y_analizar(document_id, agentes_keys, resumenes_previos, grupo):
    """Si la respuesta del panel no se puede leer, se reintenta en dos mitades."""
    mitad = max(1, len(agentes_keys) // 2)
    resultados, errores = {}, []
    for sub in (agentes_keys[:mitad], agentes_keys[mitad:]):
        if not sub:
            continue
        r, err = _analizar_grupo(document_id, sub, resumenes_previos, grupo, _reintento=True)
        resultados.update(r)
        if err:
            errores.append(err)
    return resultados, ("; ".join(errores) if errores else None)


def _analizar_grupo(document_id, agentes_keys, resumenes_previos=None, grupo=None, _reintento=False):
    grupo = grupo or {}
    if not REINTENTAR_ANALISIS:
        _reintento = True
    try:
        chunks = seleccionar_chunks_para_grupo(document_id, agentes_keys, max_chars=grupo.get("contexto_chars"))
        if not chunks:
            return {}, "Sin fragmentos"
        contexto = _construir_contexto_analisis(chunks)
        prompt_sistema = _construir_prompt_grupo(
            agentes_keys, resumenes_previos, integrador=bool(grupo.get("resumen_extendido")))
        prompt_usuario = f"Fragmentos del documento:\n\n{contexto}\n\nDevuelve el JSON del panel."
        respuesta_raw = llamar_llm_analisis(
            prompt_sistema=prompt_sistema, prompt_usuario=prompt_usuario,
            temperatura=0.1, max_tokens=grupo.get("max_tokens", MAX_TOKENS_GRUPO), json_mode=True,
        )
    except Exception as e:
        print(f"[GRUPO] Error: {e}", flush=True)
        return {}, str(e)

    try:
        data = parsear_json_respuesta(respuesta_raw)
        if not isinstance(data, dict):
            raise ValueError("la respuesta no es un objeto JSON")
    except Exception as e:
        print(f"[GRUPO] JSON ilegible ({agentes_keys}): {e}", flush=True)
        if len(agentes_keys) > 1 and not _reintento:
            return _dividir_y_analizar(document_id, agentes_keys, resumenes_previos, grupo)
        return {}, f"Respuesta ilegible: {e}"

    data_min = {str(k).lower().strip(): v for k, v in data.items()}
    resultados, faltan = {}, []
    for key in agentes_keys:
        r = data_min.get(key)
        if r is None and len(agentes_keys) == 1 and "resumen_ejecutivo" in data:
            r = data  
        if not _analisis_valido(r):
            faltan.append(key)
            continue
        r = _normalizar_resultado_analisis(r)
        resultados[key] = r
        _guardar_analisis(document_id, key, r)

    error = None
    if faltan:
        print(f"[GRUPO] Sin respuesta de: {faltan}", flush=True)
        if not _reintento:
            extra, error = _analizar_grupo(document_id, faltan, resumenes_previos, grupo, _reintento=True)
            resultados.update(extra)
            faltan = [k for k in faltan if k not in extra]
            if not faltan:
                error = None
        if faltan and not error:
            error = "Sin respuesta válida de: " + ", ".join(faltan)
    return resultados, error


def ejecutar_analisis_completo(document_id, on_progreso=None):
    memoria = {}
    completados = 0
    total = len(TODOS_LOS_AGENTES)
    for i, grupo in enumerate(GRUPOS_ANALISIS):
        agentes_a_analizar = []
        for key in grupo["agentes"]:
            existente = _obtener_analisis_guardado(document_id, key)
            if _analisis_valido(existente):
                memoria[key] = existente
                completados += 1
            else:
                agentes_a_analizar.append(key)
        if not agentes_a_analizar:
            if on_progreso: on_progreso(completados, total)
            continue
        resumenes_para_prompt = _resumenes_previos_para_grupo(document_id, i, memoria)
        print(f"[ANALISIS] {grupo['nombre']} ({i+1}/{len(GRUPOS_ANALISIS)}): {agentes_a_analizar}", flush=True)
        resultados, error = _analizar_grupo(document_id, agentes_a_analizar, resumenes_para_prompt, grupo)
        for key, r in resultados.items():
            memoria[key] = r
            completados += 1
            print(f"[ANALISIS] {key} OK (score {r.get('score')})", flush=True)
        if error:
            print(f"[ANALISIS] {grupo['nombre']}: {error}", flush=True)
        if on_progreso: on_progreso(completados, total)
        if i < len(GRUPOS_ANALISIS) - 1:
            time.sleep(PAUSA_ENTRE_GRUPOS)
    return completados, total


def _procesar_documento_bg(document_id, ruta_temporal, nombre_archivo):
    try:
        progreso_documentos[document_id]["estado"] = "Leyendo archivo"
        bloques, total_paginas = _extraer_texto_archivo(ruta_temporal, nombre_archivo)
        guardar_documento(document_id, nombre_archivo, total_paginas)
        progreso_documentos[document_id]["estado"] = "Generando chunks"
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
        items = []
        for texto, pagina in bloques:
            if not texto.strip(): continue
            for sub in splitter.split_text(texto):
                if sub.strip():
                    items.append({"content": sub.strip(), "page_number": pagina})
        total_items = len(items)
        estimado_embed = int(total_items * SEGUNDOS_POR_CHUNK_EMBEDDING) + 5
        estimado_analisis = len(GRUPOS_ANALISIS) * (SEGUNDOS_POR_GRUPO + PAUSA_ENTRE_GRUPOS)
        progreso_documentos[document_id].update({
            "fase": "embeddings", "total": total_items, "procesados": 0,
            "errores": 0, "estado": "Procesando",
            "estimado_seg": estimado_embed + estimado_analisis,
        })
        cantidad_ok = 0
        cantidad_err = 0
        inicio = time.time()
        for i in range(0, total_items, TAMANO_LOTE_EMBEDDING):
            lote = items[i:i + TAMANO_LOTE_EMBEDDING]
            textos = [it["content"] for it in lote]
            try:
                embeddings = generar_embeddings_batch(textos)
                items_con_emb = [
                    {"content": it["content"], "page_number": it["page_number"], "embedding": emb}
                    for it, emb in zip(lote, embeddings)
                ]
                guardar_chunks_batch(document_id, items_con_emb)
                cantidad_ok += len(items_con_emb)
                for _ in range(len(items_con_emb)):
                    progreso_documentos[document_id]["procesados"] = cantidad_ok + cantidad_err
                    time.sleep(0.01)
            except Exception as e:
                cantidad_err += len(lote)
                progreso_documentos[document_id]["errores"] = cantidad_err
                print(f"[EMBEDDING] Lote falló: {e}", flush=True)
            restante = max(0, estimado_embed - int(time.time() - inicio))
            progreso_documentos[document_id]["estimado_seg"] = restante + estimado_analisis
            time.sleep(PAUSA_ENTRE_LOTES)
        if cantidad_ok == 0:
            actualizar_estado(document_id, "Error")
            progreso_documentos[document_id]["estado"] = "Error"
            return
        print(f"[BG] Embeddings listos ({cantidad_ok}). Iniciando análisis...", flush=True)
        actualizar_estado(document_id, "Analizando")
        progreso_documentos[document_id].update({
            "fase": "analisis", "total": len(TODOS_LOS_AGENTES),
            "procesados": 0, "estado": "Analizando", "estimado_seg": estimado_analisis,
        })
        inicio_analisis = time.time()
        def on_progreso(completados, total):
            progreso_documentos[document_id]["procesados"] = completados
            restante = max(0, estimado_analisis - int(time.time() - inicio_analisis))
            progreso_documentos[document_id]["estimado_seg"] = restante
        completados, total_agentes = ejecutar_analisis_completo(document_id, on_progreso=on_progreso)
        if completados < total_agentes:
            print(f"[BG] Análisis incompleto: {completados}/{total_agentes} agentes", flush=True)
        estado_final = "Listo" if (cantidad_err == 0 and completados >= total_agentes) else "Listo con errores"
        actualizar_estado(document_id, estado_final)
        progreso_documentos[document_id].update({"fase": "listo", "estado": estado_final, "estimado_seg": 0})
        print(f"[BG] Documento {document_id} listo", flush=True)
    except Exception as e:
        print(f"[BG] Error: {e}", flush=True)
        traceback.print_exc()
        actualizar_estado(document_id, "Error")
        progreso_documentos[document_id]["estado"] = "Error"
    finally:
        if os.path.exists(ruta_temporal):
            try: os.remove(ruta_temporal)
            except Exception: pass


@app.get("/")
def inicio():
    return {"mensaje": "Backend Directa funcionando correctamente"}


@app.get("/api/agentes")
def api_agentes():
    return {"agentes": listar_agentes()}


@app.get("/api/progreso/{document_id}")
def consultar_progreso(document_id: str):
    if document_id not in progreso_documentos:
        return {"estado": "No encontrado", "total": 0, "procesados": 0, "fase": "desconocida"}
    return progreso_documentos[document_id]


_CACHE_CHUNKS = {}


@app.get("/api/documentos")
def listar_documentos():
    try:
        resp = (supabase.table("documentos")
                .select("id, titulo, nombre_archivo, total_paginas, estado, creado_en")
                .order("creado_en", desc=True).execute())
        documentos = resp.data or []
        conteo_an = {}
        try:
            for fila in (supabase.table("analisis").select("document_id").execute().data or []):
                conteo_an[fila["document_id"]] = conteo_an.get(fila["document_id"], 0) + 1
        except Exception:
            pass
        ahora = time.time()
        for doc in documentos:
            doc["analisis_count"] = conteo_an.get(doc["id"], 0)
            hit = _CACHE_CHUNKS.get(doc["id"])
            if hit and ahora - hit[0] < 60:
                doc["chunks"] = hit[1]
                continue
            try:
                cnt = supabase.table("document_chunks").select("id", count="exact").eq("document_id", doc["id"]).limit(1).execute()
                doc["chunks"] = cnt.count or 0
                _CACHE_CHUNKS[doc["id"]] = (ahora, doc["chunks"])
            except Exception:
                doc["chunks"] = 0
        return {"ok": True, "documentos": documentos}
    except Exception as e:
        print("[DOCUMENTOS] Error:", e, flush=True)
        return {"ok": False, "error": str(e), "documentos": []}


@app.get("/api/diagnostico/{document_id}")
def diagnostico(document_id: str):
    try:
        doc = supabase.table("documentos").select("*").eq("id", document_id).execute()
        chunks = supabase.table("document_chunks").select("content, page_number").eq("document_id", document_id).limit(10).execute()
        total_chunks = supabase.table("document_chunks").select("id", count="exact").eq("document_id", document_id).execute()
        analisis = supabase.table("analisis").select("agente, score, resultado").eq("document_id", document_id).execute()
        return {
            "ok": True,
            "documento": doc.data[0] if doc.data else None,
            "total_chunks": total_chunks.count,
            "primeros_chunks": chunks.data or [],
            "analisis": [
                {"agente": a["agente"], "score": a["score"],
                 "resumen_len": len((a["resultado"] or {}).get("resumen_ejecutivo", "")),
                 "hallazgos": len((a["resultado"] or {}).get("hallazgos", [])),
                 "senales": len((a["resultado"] or {}).get("senales", [])),
                 "preguntas": len((a["resultado"] or {}).get("preguntas", [])),
                 "recomendaciones": len((a["resultado"] or {}).get("recomendaciones", []))}
                for a in (analisis.data or [])
            ]
        }
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.delete("/api/documentos/{document_id}")
def eliminar_documento(document_id: str):
    try:
        supabase.table("document_chunks").delete().eq("document_id", document_id).execute()
        supabase.table("analisis").delete().eq("document_id", document_id).execute()
        supabase.table("documentos").delete().eq("id", document_id).execute()
        if document_id in progreso_documentos:
            del progreso_documentos[document_id]
        return {"ok": True, "mensaje": "Documento eliminado"}
    except Exception as e:
        print("[ELIMINAR] Error:", e, flush=True)
        return {"ok": False, "error": str(e)}


@app.post("/api/subir-informe/")
async def subir_informe(document_id: str = Form(...), archivo: UploadFile = File(...)):
    id_seguro = re.sub(r"[^A-Za-z0-9_.-]", "_", document_id)[:80]
    nombre_seguro = re.sub(r"[^A-Za-z0-9_.-]", "_", os.path.basename((archivo.filename or "archivo").replace("\\", "/")))[:120]
    ruta_temporal = f"temp_{id_seguro}_{nombre_seguro}"
    try:
        progreso_documentos[document_id] = {
            "fase": "iniciando", "total": 0, "procesados": 0,
            "errores": 0, "estado": "Iniciando", "estimado_seg": 0,
        }
        with open(ruta_temporal, "wb") as buffer:
            shutil.copyfileobj(archivo.file, buffer)
        thread = threading.Thread(
            target=_procesar_documento_bg,
            args=(document_id, ruta_temporal, archivo.filename),
            daemon=True,
        )
        thread.start()
        return {"ok": True, "mensaje": "Documento recibido.", "document_id": document_id}
    except Exception as e:
        traceback.print_exc()
        if os.path.exists(ruta_temporal):
            os.remove(ruta_temporal)
        return {"ok": False, "error": str(e)}

def buscar_chunks(document_id, pregunta, match_threshold=0.05, match_count=8):
    emb = generar_embedding_pregunta(pregunta)
    rpc = supabase.rpc("match_document_chunks", {
        "query_embedding": emb, "match_threshold": match_threshold,
        "match_count": match_count, "p_document_id": document_id,
    }).execute()
    return rpc.data or []


def _construir_contexto_analisis_para_chat(analisis, agente_foco=None):
    if not analisis: return ""
    partes = []
    if agente_foco and agente_foco in analisis:
        r = analisis[agente_foco]
        partes.append(f"[{agente_foco.upper()} - FOCO PRINCIPAL]")
        if r.get("resumen_ejecutivo"):
            partes.append(f"Resumen: {r['resumen_ejecutivo']}")
        for h in (r.get("hallazgos") or [])[:5]:
            if isinstance(h, dict):
                partes.append(f"- Hallazgo: {h.get('titulo','')}. {h.get('descripcion','')}")
        for s in (r.get("senales") or [])[:5]:
            if isinstance(s, dict):
                partes.append(f"- Señal ({s.get('severidad','')}): {s.get('titulo','')}. {s.get('descripcion','')}")
    partes.append("\n[RESÚMENES DE OTROS AGENTES]")
    for ag, r in analisis.items():
        if ag == agente_foco: continue
        if isinstance(r, dict) and r.get("resumen_ejecutivo"):
            partes.append(f"[{ag.upper()}] {r['resumen_ejecutivo']}")
    return "\n".join(partes)


def _parsear_historial(historial, max_turnos=4, max_chars=700):
    """Turnos previos del chat (los manda el frontend como JSON). Tolerante: si viene mal, se ignora."""
    if not historial:
        return []
    try:
        datos = json.loads(historial)
    except Exception:
        return []
    out = []
    if isinstance(datos, list):
        for t in datos[-max_turnos:]:
            if isinstance(t, dict):
                p = str(t.get("pregunta") or "").strip()[:max_chars]
                r = str(t.get("respuesta") or "").strip()[:max_chars]
                if p:
                    out.append({"pregunta": p, "respuesta": r})
    return out


def _pregunta_para_busqueda(pregunta, historial_lista):
    """Una repregunta corta ('¿y por qué?') no sirve para buscar fragmentos: se le suma la pregunta anterior."""
    pregunta = re.sub(r"^\s*\[[^\]]{1,30}\]\s*", "", pregunta)
    pregunta = re.sub(r"^\s*Expl[ií]came con las cifras del documento\s*:\s*", "", pregunta, flags=re.I)
    if historial_lista and len(pregunta.split()) <= 6:
        return f"{historial_lista[-1]['pregunta']} {pregunta}"
    return pregunta


def construir_prompt_chat(config, chunks, pregunta, analisis_previos, agente_foco=None, historial_lista=None, nombre_doc=""):
    contexto_chunks = "\n\n".join(f"[Página {c['page_number']}]\n{c['content']}" for c in chunks)
    contexto_analisis = _construir_contexto_analisis_para_chat(analisis_previos, agente_foco)
    prompt_sistema = (
        f"{config['rol_chat']} "
        "Eres un asistente ejecutivo que conversa con el Gerente General. Responde en español, claro y directo. "
        "IMPORTANTE: 'Directa' es la plataforma que usas, NO es la empresa del documento; el nombre de la empresa "
        "y el período salen de los fragmentos. "
        "Empieza SIEMPRE con la respuesta directa en la primera frase, sin preámbulos. "
        "Si preguntan de qué trata el documento o piden un resumen general, empieza con 'El documento trata de…' "
        "(tipo de documento, empresa, período y propósito), sigue con 3 a 5 puntos clave con cifras en viñetas "
        "y cierra con una línea sobre lo más importante para el directorio. "
        "No uses tablas salvo que el usuario las pida. No escribas etiquetas como [ANÁLISIS PREVIOS…] ni "
        "[Fragmentos pág. N]: cita la página al final de la frase como (pág. N). "
        "Usa un máximo de 250 palabras salvo que pidan más detalle. "
        "Si piden cifras, entrega las cifras exactas de los fragmentos (incluidos los cuadros, donde las columnas "
        "vienen separadas por ' | ') con su unidad y el año; compara fila por fila y no te quedes solo con una frase general. "
        "No uses símbolos de cita especiales (【 】 ■) ni escribas 'ANÁLISIS PREVIOS' en la respuesta. "
        "Los datos y las explicaciones salen de los FRAGMENTOS DEL DOCUMENTO; los ANÁLISIS PREVIOS "
        "solo te orientan sobre qué mirar. Son resúmenes automáticos y pueden equivocarse: si un "
        "fragmento muestra una explicación o una cifra que el análisis dice que no existe, manda "
        "el fragmento y aclara que el análisis previo no la consideró. "
        "NO presentes los ANÁLISIS PREVIOS como cita textual del documento; solo pueden ir entre "
        "comillas frases que aparezcan literalmente en los FRAGMENTOS. "
        "Cuando cites una página usa EXACTAMENTE el número de [Página N] del fragmento; no uses el "
        "número de una sección ni de un cuadro como si fuera una página. "
        "Si algo no consta en los fragmentos, dilo en una frase ('En los fragmentos revisados no "
        "aparece…') y, si ayuda, di qué sí consta. NUNCA pidas la página ni preguntes dónde está. "
        "Si el usuario te corrige, acepta sin discutir."
    )
    bloque_hist = ""
    if historial_lista:
        bloque_hist = "CONVERSACIÓN PREVIA (para entender repreguntas; los datos salen del documento):\n" + "\n".join(
            f"P: {t['pregunta']}\nR: {t['respuesta']}" for t in historial_lista) + "\n\n"
    prompt_usuario = (
        (f"NOMBRE DEL ARCHIVO: {nombre_doc}\n\n" if nombre_doc else "") +
        f"ANÁLISIS PREVIOS (fuente principal):\n{contexto_analisis}\n\n"
        f"FRAGMENTOS DEL DOCUMENTO (respaldo):\n{contexto_chunks}\n\n"
        f"{bloque_hist}"
        f"Pregunta: {pregunta}\n\nRespuesta:"
    )
    return prompt_sistema, prompt_usuario


_RE_PIDE_PAGINA = re.compile(
    r"[^.!?\n]*(?:¿[^?\n]*)?(?:ind[ií]c(?:a|ame|arme|as|ar)|dime|dec[ií]r(?:me)?|se[ñn]ala|me\s+puedes?\s+(?:decir|indicar)|"
    r"podr[ií]as?\s+(?:decir|indicar))[^?\n]*p[áa]gina[^?\n]*\?", re.IGNORECASE)


def _quitar_pedido_de_pagina(texto):
    """Red de seguridad: el modelo a veces cierra con '¿Me indicas la página?'. Se elimina esa frase."""
    if not texto:
        return texto
    limpio = _RE_PIDE_PAGINA.sub("", texto)
    # Marcas de cita que algunos modelos agregan (【…】, ■ANÁLISIS PREVIOS■): no se ven bien en pantalla ni en el PDF.
    limpio = re.sub(r"[【\[]\s*(?:■\s*)?AN[ÁA]LISIS PREVIOS?[^】\]]{0,60}[】\]]", "", limpio, flags=re.I)
    limpio = re.sub(r"【[^】]{0,200}】", "", limpio)
    limpio = re.sub(r"■\s*AN[ÁA]LISIS PREVIOS?\s*■?", "", limpio, flags=re.I)
    limpio = limpio.replace("■", "").replace("【", "(").replace("】", ")")
    limpio = re.sub(r"\s*\(\s*p[áa]gs?\.?\s*[-‐‑–—?]*\s*\)", "", limpio, flags=re.I)   # "(pág. -)" sin número
    limpio = re.sub(r" +([.,;:])", r"\1", limpio)
    limpio = re.sub(r"[ \t]+\n", "\n", limpio)
    return re.sub(r"\n{3,}", "\n\n", limpio).strip() or texto


_PALABRAS_VACIAS = {
    "cual", "cuales", "como", "cuando", "donde", "quien", "porque", "existe", "existen", "tiene", "tienen",
    "dispone", "disponen", "documento", "informe", "principal", "causa", "plan", "accion", "revertirla",
    "permita", "comparar", "variaciones", "observadas", "sobre", "entre", "desde", "hasta", "segun",
    "resultado", "resultados", "ejercicio", "financiero", "financieros", "compania", "empresa", "importante",
    "relevante", "relevantes", "informacion", "principales",
}


def _terminos_clave(pregunta, maximo=2):
    """Siglas (EBITDA, HLAG, CAPEX…) primero y, si no hay, palabras largas poco comunes."""
    siglas = re.findall(r"\b[A-ZÁÉÍÓÚÑ]{3,}[A-Z0-9]*\b", pregunta or "")
    out = []
    for t in siglas:
        if t.lower() not in out:
            out.append(t.lower())
    if len(out) < maximo:
        for w in re.findall(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{9,}", pregunta or ""):
            wl = _normalizar_texto(w)
            if wl not in _PALABRAS_VACIAS and wl not in out:
                out.append(wl)
    return out[:maximo]


def ampliar_chunks(document_id, chunks, pregunta, max_total=10):
    """Suma al resultado de la búsqueda por significado (a) fragmentos que nombran el término
    clave de la pregunta (una sigla como EBITDA) y (b) el fragmento que sigue al mejor resultado,
    porque la explicación de una tabla suele estar justo después. Sin llamadas a la IA."""
    try:
        vistos = {c.get("id") for c in chunks if c.get("id") is not None}
        extra = []
        for t in _terminos_clave(pregunta):
            resp = (supabase.table("document_chunks").select("id, content, page_number")
                    .eq("document_id", document_id).ilike("content", f"%{t}%").limit(25).execute())
            cand = [c for c in (resp.data or []) if c.get("id") not in vistos]
            cand.sort(key=lambda c: -(_bonus_chunk(_normalizar_texto(c.get("content") or ""))))
            for c in cand[:1]:
                extra.append(c); vistos.add(c.get("id"))
        if chunks:
            mejores = sorted(chunks, key=lambda c: -(c.get("similarity") or 0))[:2]
            for c in mejores:
                cid = c.get("id")
                if cid is None:   # la búsqueda por significado puede no devolver el id: se ubica por su texto
                    r0 = (supabase.table("document_chunks").select("id").eq("document_id", document_id)
                          .eq("content", c.get("content") or "").limit(1).execute())
                    cid = (r0.data or [{}])[0].get("id")
                if cid is None:
                    continue
                resp = (supabase.table("document_chunks").select("id, content, page_number")
                        .eq("document_id", document_id).gt("id", cid).order("id").limit(1).execute())
                for n in (resp.data or []):
                    if n.get("id") not in vistos:
                        extra.append(n); vistos.add(n.get("id"))
        for c in extra:
            c.setdefault("similarity", 0.0)
        return (list(chunks) + extra)[:max_total]
    except Exception as e:
        print("[CHAT] No se pudo ampliar la búsqueda:", e, flush=True)
        return chunks


_RE_PREGUNTA_GENERAL = re.compile(
    r"(de\s+qu[eé]\s+(trata|se\s+trata|habla|es)|qu[eé]\s+es\s+(este|el)\s+(documento|informe|archivo|reporte)|"
    r"resum(e|en|ir|ime)|resumen\s+general|descr[ií]b(e|eme)\s+(el|este)\s+(documento|informe)|"
    r"de\s+qu[eé]\s+(documento|informe)|qu[eé]\s+contiene|visi[oó]n\s+general|panorama\s+general)", re.IGNORECASE)


def _es_pregunta_general(pregunta):
    return bool(_RE_PREGUNTA_GENERAL.search(pregunta or ""))


def _agregar_inicio_documento(document_id, chunks, pregunta, n=4):
    """Para 'de qué trata el documento' se suman los primeros fragmentos (portada, índice, mensaje inicial)
    y el nombre del archivo; la búsqueda por significado sola suele traer páginas sueltas."""
    if not _es_pregunta_general(pregunta):
        return chunks
    try:
        resp = (supabase.table("document_chunks").select("id, content, page_number")
                .eq("document_id", document_id).order("id").limit(n).execute())
        vistos = {c.get("id") for c in chunks if c.get("id") is not None}
        extra = [c for c in (resp.data or []) if c.get("id") not in vistos]
        for c in extra:
            c["similarity"] = 0.0
        return extra + list(chunks)
    except Exception as e:
        print("[CHAT] No se pudo leer el inicio del documento:", e, flush=True)
        return chunks


def _nombre_documento(document_id):
    try:
        r = supabase.table("documentos").select("titulo, nombre_archivo").eq("id", document_id).limit(1).execute()
        d = (r.data or [{}])[0]
        return d.get("titulo") or d.get("nombre_archivo") or ""
    except Exception:
        return ""


def extraer_fuentes(chunks):
    return [{"pagina": c["page_number"], "preview": c["content"][:200] + "...", "similarity": round(c["similarity"], 3)} for c in chunks]


@app.post("/api/consultar/")
async def consultar(pregunta: str = Form(...), document_id: str = Form(...), agente: str = Form("general"),
                    historial: str = Form("")):
    try:
        config = obtener_agente(agente)
        hist = _parsear_historial(historial)
        chunks = _agregar_inicio_documento(document_id, ampliar_chunks(document_id, buscar_chunks(document_id, _pregunta_para_busqueda(pregunta, hist)), pregunta), pregunta)
        analisis = _obtener_todos_los_analisis(document_id)
        if not chunks and not analisis:
            return {"ok": True, "respuesta": "No encontré información.", "fuentes": [],
                    "agente": config["nombre"], "agente_key": agente}
        agente_foco = agente if agente in analisis else None
        prompt_sistema, prompt_usuario = construir_prompt_chat(config, chunks, pregunta, analisis, agente_foco, hist, _nombre_documento(document_id) if _es_pregunta_general(pregunta) else "")
        respuesta = _quitar_pedido_de_pagina(llamar_llm_chat(prompt_sistema=prompt_sistema, prompt_usuario=prompt_usuario, temperatura=0.2, max_tokens=1000))
        return {"ok": True, "respuesta": respuesta, "fuentes": extraer_fuentes(chunks),
                "agente": config["nombre"], "agente_key": agente}
    except Exception as e:
        print("[CONSULTAR] Error:", e, flush=True)
        return {"ok": False, "error": str(e)}


@app.post("/api/consultar-stream/")
async def consultar_stream(pregunta: str = Form(...), document_id: str = Form(...), agente: str = Form("general"),
                           historial: str = Form("")):
    config = obtener_agente(agente)
    print(f"[STREAM] Agente: {config['nombre']}", flush=True)
    hist = _parsear_historial(historial)
    try:
        chunks = _agregar_inicio_documento(document_id, ampliar_chunks(document_id, buscar_chunks(document_id, _pregunta_para_busqueda(pregunta, hist)), pregunta), pregunta)
        analisis = _obtener_todos_los_analisis(document_id)
    except Exception as e:
        async def err():
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
            yield "data: [DONE]\n\n"
        return StreamingResponse(err(), media_type="text/event-stream")
    if not chunks and not analisis:
        async def vacio():
            yield f"data: {json.dumps({'chunk': 'No encontré información.'})}\n\n"
            yield f"data: {json.dumps({'fuentes': [], 'agente': config['nombre']})}\n\n"
            yield "data: [DONE]\n\n"
        return StreamingResponse(vacio(), media_type="text/event-stream")
    agente_foco = agente if agente in analisis else None
    prompt_sistema, prompt_usuario = construir_prompt_chat(config, chunks, pregunta, analisis, agente_foco, hist, _nombre_documento(document_id) if _es_pregunta_general(pregunta) else "")
    fuentes = extraer_fuentes(chunks)
    async def stream_llm():
        try:
            respuesta = _quitar_pedido_de_pagina(llamar_llm_chat(prompt_sistema=prompt_sistema, prompt_usuario=prompt_usuario, temperatura=0.2, max_tokens=1000))
            if respuesta:
                for i in range(0, len(respuesta), 15):
                    yield f"data: {json.dumps({'chunk': respuesta[i:i+15]})}\n\n"
            yield f"data: {json.dumps({'fuentes': fuentes, 'agente': config['nombre']})}\n\n"
            yield "data: [DONE]\n\n"
        except Exception as e:
            print("[STREAM] Error:", e, flush=True)
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
            yield "data: [DONE]\n\n"
    return StreamingResponse(stream_llm(), media_type="text/event-stream")


# ============================================================================
# SUGERENCIAS
# ============================================================================
def _titulo_corto(item, limite=90):
    t = (item.get("titulo") if isinstance(item, dict) else str(item)) or ""
    t = " ".join(t.split()).strip(" .:")
    return t[:limite].rstrip()


def _sugerencias_de_agente(r, etiqueta=""):
    """Preguntas que el documento SÍ puede contestar: salen de lo que el agente encontró (hallazgos) y de
    sus alertas (para comprobarlas contra el texto). Las 'preguntas probables del directorio' no se usan
    aquí: son huecos del documento, así que el chat no podría responderlas."""
    out = []
    for h in (r.get("hallazgos") or [])[:3]:
        t = _titulo_corto(h)
        if t:
            out.append(f"{etiqueta}Explícame con las cifras del documento: {t}")
    for sn in (r.get("senales") or [])[:2]:
        t = _titulo_corto(sn)
        if t:
            out.append(f"{etiqueta}Comprueba en el documento esta alerta: {t}")
    return out


@app.post("/api/sugerencias/")
async def sugerencias(document_id: str = Form(...), agente: str = Form("general")):
    try:
        analisis = _obtener_todos_los_analisis(document_id)
        if agente in analisis:
            sug = _sugerencias_de_agente(analisis[agente])
            if sug:
                return {"ok": True, "sugerencias": sug[:5], "fuente": "analisis_agente"}
        if agente == "general" and analisis:
            todas = []
            for ag, r in analisis.items():
                if isinstance(r, dict):
                    todas.extend(_sugerencias_de_agente(r, f"[{ag.upper()}] ")[:1])
            if todas:
                return {"ok": True, "sugerencias": todas[:6], "fuente": "analisis_combinado"}
        return {"ok": True, "sugerencias": ["¿De qué trata el documento?", "¿Qué cifras clave aparecen?", "¿Qué riesgos se mencionan?", "¿Qué resultado tuvo el ejercicio?"], "fuente": "fallback"}
    except Exception as e:
        print("[SUGERENCIAS] Error:", e, flush=True)
        return {"ok": False, "error": str(e)}


# ============================================================================
# ANÁLISIS MANUAL
# ============================================================================
@app.post("/api/analizar-iniciar/")
async def analizar_iniciar(document_id: str = Form(...), agentes: str = Form(...)):
    try:
        agentes_list = json.loads(agentes)
        if not isinstance(agentes_list, list) or not agentes_list:
            return {"ok": False, "error": "Lista vacía"}
        agentes_list = [k for k in agentes_list if k in TODOS_LOS_AGENTES]
        if not agentes_list:
            return {"ok": False, "error": "Ningún agente válido en la lista"}
        # Evita pagar dos veces el mismo análisis (doble clic, otra pestaña, documento aún procesándose)
        if progreso_documentos.get(document_id, {}).get("fase") in ("iniciando", "embeddings", "analisis"):
            return {"ok": False, "error": "El documento todavía se está procesando."}
        for j in jobs_analisis.values():
            if j["document_id"] == document_id and j.get("corriendo"):
                return {"ok": False, "error": "Ya hay un análisis en curso para este documento.", "job_id": j["job_id"]}
        job_id = str(uuid.uuid4())
        n_grupos = sum(1 for g in GRUPOS_ANALISIS if any(k in agentes_list for k in g["agentes"]))
        estimado_seg = n_grupos * (SEGUNDOS_POR_GRUPO + PAUSA_ENTRE_GRUPOS)
        jobs_analisis[job_id] = {
            "job_id": job_id, "document_id": document_id, "agentes": agentes_list,
            "resultados": {}, "corriendo": True, "detenido": False, "pausado": False,
            "iniciado_en": datetime.now().isoformat(), "estimado_seg": estimado_seg,
            "inicio_ts": time.time(), "errores": [],
        }
        def worker():
            job = jobs_analisis[job_id]
            memoria = {}
            try:
                for i, grupo in enumerate(GRUPOS_ANALISIS):
                    if job.get("detenido"): break
                    while job.get("pausado") and not job.get("detenido"):
                        time.sleep(1)
                    if job.get("detenido"): break
                    pendientes = [k for k in grupo["agentes"] if k in agentes_list]
                    if not pendientes: continue
                    a_analizar = []
                    for key in pendientes:
                        # Nunca se repite un análisis ya guardado (ahorra IA)
                        existente = _obtener_analisis_guardado(document_id, key)
                        if _analisis_valido(existente):
                            job["resultados"][key] = existente
                            memoria[key] = existente
                        else:
                            a_analizar.append(key)
                    if not a_analizar: continue
                    resumenes = _resumenes_previos_para_grupo(document_id, i, memoria)
                    resultados, error = _analizar_grupo(document_id, a_analizar, resumenes, grupo)
                    for key, r in resultados.items():
                        job["resultados"][key] = r
                        memoria[key] = r
                    if error:
                        job.setdefault("errores", []).append(f"{grupo['nombre']}: {error}")
                    quedan = any(k in agentes_list for g in GRUPOS_ANALISIS[i + 1:] for k in g["agentes"])
                    if quedan and not job.get("detenido"):
                        time.sleep(PAUSA_ENTRE_GRUPOS)
            except Exception as e:
                traceback.print_exc()
                job.setdefault("errores", []).append(str(e))
            finally:
                job["corriendo"] = False
        threading.Thread(target=worker, daemon=True).start()
        return {"ok": True, "job_id": job_id, "estimado_seg": estimado_seg}
    except Exception as e:
        print("[ANALIZAR-INICIAR] Error:", e, flush=True)
        return {"ok": False, "error": str(e)}


@app.get("/api/analizar-estado/{job_id}")
def analizar_estado(job_id: str):
    if job_id not in jobs_analisis:
        return {"ok": False, "error": "Job no encontrado"}
    job = jobs_analisis[job_id]
    restante = 0
    if job["corriendo"] and not job.get("pausado"):
        restante = max(0, int(job.get("estimado_seg", 0) - (time.time() - job.get("inicio_ts", time.time()))))
    return {"ok": True, "job_id": job_id, "corriendo": job["corriendo"], "detenido": job["detenido"],
            "pausado": job.get("pausado", False), "agentes": job["agentes"],
            "resultados": job["resultados"], "completados": len(job["resultados"]),
            "total": len(job["agentes"]), "estimado_restante_seg": restante,
            "errores": job.get("errores", [])}


@app.post("/api/analizar-detener/{job_id}")
def analizar_detener(job_id: str):
    if job_id not in jobs_analisis: return {"ok": False, "error": "Job no encontrado"}
    jobs_analisis[job_id]["detenido"] = True
    return {"ok": True}


@app.post("/api/analizar-pausar/{job_id}")
def analizar_pausar(job_id: str):
    if job_id not in jobs_analisis: return {"ok": False, "error": "Job no encontrado"}
    jobs_analisis[job_id]["pausado"] = True
    return {"ok": True}


@app.post("/api/analizar-reanudar/{job_id}")
def analizar_reanudar(job_id: str):
    if job_id not in jobs_analisis: return {"ok": False, "error": "Job no encontrado"}
    jobs_analisis[job_id]["pausado"] = False
    return {"ok": True}


@app.get("/api/analisis-completos/{document_id}")
def analisis_completos(document_id: str):
    try:
        resp = (supabase.table("analisis").select("agente, resultado, score")
                .eq("document_id", document_id).execute())
        data = resp.data or []
        return {"ok": True, "analisis": {row["agente"]: row["resultado"] for row in data}, "count": len(data)}
    except Exception as e:
        return {"ok": False, "error": str(e), "analisis": {}}


# ============================================================================
# INFORME
# ============================================================================
SEVERIDAD_ORDEN = {"crítica": 0, "critica": 0, "alta": 1, "media": 2, "baja": 3}


def _score_label(s):
    # Con el score por señales, una sola señal alta ya baja de 90: no debe verse como "sin observaciones"
    if s >= 90: return "Sin observaciones", "#2F6B4F"
    if s >= 75: return "Menor atención", "#7cb342"
    if s >= 55: return "Requiere revisión", "#A67C00"
    if s >= 40: return "Observaciones serias", "#C2571A"
    return "Observaciones críticas", "#8E1B2E"


def _severidad_color(sev):
    sev = (sev or "").lower()
    if "crít" in sev or "crit" in sev: return "#8E1B2E"
    if "alt" in sev: return "#C2571A"
    if "med" in sev: return "#A67C00"
    if "baj" in sev: return "#2F6B4F"
    return "#888"


def _es_integrador(key):
    try:
        return obtener_agente(key).get("tipo") == "integrador"
    except Exception:
        return False


def _repartir(items, max_total, max_por_agente=None):
    """Reparte la selección entre agentes (round-robin) para que uno solo no cope la lista.
    `items`: lista de (agente, valor) ya ordenada por prioridad. Mantiene el orden relativo."""
    por_agente, orden = {}, []
    for ag, v in items:
        if ag not in por_agente:
            por_agente[ag] = []
            orden.append(ag)
        por_agente[ag].append(v)
    out, ronda = [], 0
    while len(out) < max_total:
        avanzo = False
        for ag in orden:
            if max_por_agente is not None and ronda >= max_por_agente:
                continue
            if ronda < len(por_agente[ag]):
                out.append((ag, por_agente[ag][ronda]))
                avanzo = True
                if len(out) >= max_total:
                    break
        if not avanzo:
            break
        ronda += 1
    return out


def _nombre_corto_agente(key):
    try:
        nombre = obtener_agente(key).get("nombre", key)
    except Exception:
        nombre = key
    nombre = re.sub(r"^Agente\s+", "", nombre)
    nombre = re.sub(r"\s*\(.*?\)\s*$", "", nombre)
    return nombre or key


def _construir_resumen_general(items):
    analisis_por_agente = {}
    for it in items:
        if it.get("tipo") == "analisis":
            key = it.get("agente_key", "general")
            r = it.get("resultado", {})
            if isinstance(r, dict):
                analisis_por_agente[key] = r
    if not analisis_por_agente:
        return "No hay análisis disponibles."
    con_info = []
    for key, r in analisis_por_agente.items():
        if len(r.get("hallazgos", [])) > 0 or len(r.get("senales", [])) > 0:
            con_info.append(key)
    total = len(analisis_por_agente)
    n_con = len(con_info)
    partes = []
    if n_con <= 2:
        if n_con == 0:
            frase = f"Ninguno de los {total} agentes especializados identificó contenido sustantivo en el documento."
        elif n_con == 1:
            frase = f"Solo 1 de los {total} agentes especializados identificó contenido relevante para su área."
        else:
            frase = f"Solo {n_con} de los {total} agentes especializados identificaron contenido relevante."
        partes.append(frase + " El documento no contiene la información necesaria para una preparación integral de sesión de directorio.")
        return "\n\n".join(partes)
    if n_con == 1:
        frase = f"El documento fue evaluado por {total} agentes especializados. 1 de ellos identificó hallazgos relevantes para su área."
    else:
        frase = f"El documento fue evaluado por {total} agentes especializados. {n_con} de ellos identificaron hallazgos relevantes."
    partes.append(frase)

    # Datos reales del análisis: conteo por severidad, señales más graves y áreas con menor puntaje.
    conteo = {"crítica": 0, "alta": 0, "media": 0, "baja": 0}
    graves = []
    n_transversales = 0
    for key in con_info:
        for s in (analisis_por_agente[key].get("senales") or []):
            if not isinstance(s, dict):
                continue
            sev = _norm_sev(s.get("severidad"))
            conteo[sev] += 1
            if _es_integrador(key):
                n_transversales += 1
            if sev in ("crítica", "alta") and str(s.get("titulo") or "").strip():
                graves.append((_ORDEN_SEV[sev], key, str(s["titulo"]).strip(), sev))
    n_senales = sum(conteo.values())
    if n_senales:
        plural = {"crítica": "críticas", "alta": "altas", "media": "medias", "baja": "bajas"}
        detalle = ", ".join(f"{conteo[k]} {plural[k]}" for k in ("crítica", "alta", "media", "baja") if conteo[k])
        extra = (f" {n_transversales} de ellas son transversales (Riesgos y Estrategia cruzan hechos que otros "
                 "agentes ya reportaron, por lo que pueden repetir un mismo tema).") if n_transversales else ""
        partes.append(f"En total se levantaron {n_senales} señales: {detalle}.{extra}")
    # Primero las de los agentes especialistas; las transversales (que repiten temas) después, a igual severidad
    graves.sort(key=lambda g: (g[0], _es_integrador(g[1])))
    if graves:
        partes.append("Señales de mayor severidad: " + "; ".join(
            f"{_nombre_corto_agente(k)}: {t} ({sev})" for _, k, t, sev in graves[:5]) + ".")
    con_score = []
    for key in con_info:
        try:
            con_score.append((int(analisis_por_agente[key].get("score", 0)), key))
        except (TypeError, ValueError):
            pass
    bajos = [(sc, k) for sc, k in sorted(con_score) if sc < 70][:3]
    if bajos:
        partes.append("Áreas que requieren más atención según su puntaje: " + ", ".join(
            f"{_nombre_corto_agente(k)} ({sc}/100)" for sc, k in bajos) + ".")
    partes.append("A continuación se presenta la síntesis de los aspectos más importantes.")
    return "\n\n".join(partes)


@app.post("/api/informe/")
async def generar_informe(document_id: str = Form(...), payload: str = Form(...)):
    try:
        data = json.loads(payload)
        items = data.get("items", [])
        doc_info = data.get("documento", {})
        por_agente = {}
        for it in items:
            key = it.get("agente_key") or it.get("agente") or "general"
            nombre = it.get("agente") or "Asistente General"
            if key not in por_agente:
                por_agente[key] = {"nombre": nombre, "chats": [], "analisis": []}
            if it.get("tipo") == "chat":
                por_agente[key]["chats"].append(it)
            elif it.get("tipo") == "analisis":
                por_agente[key]["analisis"].append(it)
        resumen_general = _construir_resumen_general(items)
        html = _construir_html_informe(doc_info, resumen_general, por_agente)

        # Barrido FINAL: convierte tildes a entidades, limpia caracteres raros
        html = _entidades_para_pdf(html)

        pdf_buffer = io.BytesIO()
        pisa_status = pisa.CreatePDF(src=html, dest=pdf_buffer, encoding="utf-8")
        if pisa_status.err:
            raise Exception("Error generando PDF")
        pdf_buffer.seek(0)
        nombre_pdf = f"Directa_Informe_{datetime.now().strftime('%Y-%m-%d')}.pdf"
        return StreamingResponse(pdf_buffer, media_type="application/pdf",
                                 headers={"Content-Disposition": f'attachment; filename="{nombre_pdf}"'})
    except Exception as e:
        print("[INFORME] Error:", e, flush=True)
        traceback.print_exc()
        return HTMLResponse(content=f"<h1>Error</h1><pre>{traceback.format_exc()}</pre>", status_code=500)


def _construir_html_informe(doc_info, resumen_general, por_agente):
    fecha = datetime.now().strftime("%d/%m/%Y %H:%M")
    nombre_doc = _esc_html(doc_info.get("nombre", "Documento"))
    paginas = doc_info.get("paginas", "?")

    todos_analisis = []
    for key, data in por_agente.items():
        for a in data["analisis"]:
            r = a.get("resultado", {})
            if isinstance(r, dict):
                todos_analisis.append({"agente": key, "nombre": data["nombre"], "resultado": r})

    # FIX: el score global se calcula SOLO sobre agentes que evaluaron el
    # documento (con hallazgos o señales). Los "vacíos" (score 90 sin items)
    # no aportan información sobre la preparación real y distorsionarían el promedio.
    analisis_con_info = []
    for a in todos_analisis:
        r = a["resultado"]
        tiene_info = bool(r.get("hallazgos") or r.get("senales"))
        if tiene_info:
            analisis_con_info.append(a)

    n_evaluados = len(analisis_con_info)
    n_total = len(todos_analisis)

    if analisis_con_info:
        scores = [a["resultado"].get("score", 0) for a in analisis_con_info]
        score_global = round(sum(scores) / len(scores))
    else:
        score_global = 0

    score_global_label, score_global_color = _score_label(score_global)

    senales_criticas = []
    for a in todos_analisis:
        for s in (a["resultado"].get("senales") or []):
            if isinstance(s, dict):
                sev = (s.get("severidad") or "").lower()
                if sev in ("crítica", "critica", "alta"):
                    senales_criticas.append({
                        "agente": a["agente"], "nombre_agente": a["nombre"],
                        "titulo": s.get("titulo", ""), "descripcion": s.get("descripcion", ""),
                        "severidad": sev,
                    })
    # Primero severidad; a igual severidad, los agentes especialistas antes que los transversales
    # (Riesgos/Estrategia). Máximo 2 por agente para que uno solo no ocupe la lista.
    senales_criticas.sort(key=lambda x: (SEVERIDAD_ORDEN.get(x["severidad"], 99), _es_integrador(x["agente"])))
    top_senales = [s for _, s in _repartir([(s["agente"], s) for s in senales_criticas], 10, max_por_agente=2)]
    top_senales.sort(key=lambda x: (SEVERIDAD_ORDEN.get(x["severidad"], 99), _es_integrador(x["agente"])))

    # Una pregunta por agente (empezando por los de menor puntaje), luego una segunda ronda, hasta 12.
    def _score_ord(a):
        try:
            return int(a["resultado"].get("score", 100))
        except (TypeError, ValueError):
            return 100
    todas_preguntas = []
    for a in sorted(todos_analisis, key=_score_ord):
        for p in (a["resultado"].get("preguntas") or []):
            txt = p.get("pregunta", "") if isinstance(p, dict) else str(p)
            if txt.strip():
                todas_preguntas.append((a["nombre"], txt.strip()))
    top_preguntas = _repartir(todas_preguntas, 12)

    from modulos_base import tabla_pdf as _tabla, color_texto as _color
    _titulo_seccion = '<p style="font-size:8pt;font-weight:bold;color:#1B2F4B;margin:8pt 0 2pt 0;">{}</p>'

    def _sev_celda(sev):
        return _color(str(sev).upper(), _severidad_color(sev))

    def bloque_senales_html(senales):
        if not senales: return ""
        filas = [[_sev_celda(s["severidad"]),
                  f'<b>{_esc_html(s["titulo"])}</b><br/><span style="color:#5F6877;">{_esc_html(s["descripcion"])}</span>',
                  _esc_html(s["nombre_agente"])] for s in senales]
        return _tabla(["Severidad", "Señal detectada", "Área"], filas, ["12%", "68%", "20%"])

    def bloque_preguntas_html(preguntas):
        if not preguntas:
            return "<p style='color:#5F6877;'>No se generaron preguntas.</p>"
        filas = [[str(idx), _esc_html(p), _esc_html(agente_nombre)] for idx, (agente_nombre, p) in enumerate(preguntas, 1)]
        return _tabla(["N°", "Pregunta probable", "Área"], filas, ["6%", "74%", "20%"])

    def _barra_puntaje(score, color):
        try:
            s = max(0, min(100, int(score)))
        except (TypeError, ValueError):
            s = 0
        if s >= 100:
            return f'<table width="100%" cellpadding="0" cellspacing="0"><tr><td bgcolor="{color}" style="font-size:2pt;line-height:4pt;">&nbsp;</td></tr></table>'
        if s <= 0:
            return '<table width="100%" cellpadding="0" cellspacing="0"><tr><td bgcolor="#E6EAF0" style="font-size:2pt;line-height:4pt;">&nbsp;</td></tr></table>'
        return (f'<table width="100%" cellpadding="0" cellspacing="0"><tr><td width="{s}%" bgcolor="{color}" style="font-size:2pt;line-height:4pt;">&nbsp;</td>'
                f'<td bgcolor="#E6EAF0" style="font-size:2pt;line-height:4pt;">&nbsp;</td></tr></table>')

    def _lista(items, campo):
        textos = []
        for x in (items or []):
            txt = x if isinstance(x, str) else (x.get(campo, "") if isinstance(x, dict) else "")
            if str(txt).strip():
                textos.append(str(txt).strip())
        if not textos:
            return ""
        return "<ul>" + "".join(f"<li>{_esc_html(t)}</li>" for t in textos) + "</ul>"

    def render_analisis_agente(r, score):
        sc_label, sc_color = _score_label(score)
        resumen_html = _render_markdown_a_html(r.get("resumen_ejecutivo", "Sin resumen ejecutivo."))
        out = (f'<table width="100%" cellpadding="0" cellspacing="0"><tr>'
               f'<td width="14%" style="font-size:16pt;font-weight:bold;color:{sc_color};">{score}<span style="font-size:8pt;color:#5F6877;font-weight:normal;"> /100</span></td>'
               f'<td width="26%" style="font-size:9pt;font-weight:bold;color:{sc_color};">{sc_label}</td>'
               f'<td>{_barra_puntaje(score, sc_color)}</td></tr></table>')
        out += f'<div style="color:#2B2F36;font-size:9.5pt;line-height:1.45;margin-top:4pt;">{resumen_html}</div>'
        hallazgos = []
        for h in (r.get("hallazgos") or []):
            if isinstance(h, str): h = {"titulo": h, "descripcion": ""}
            if isinstance(h, dict):
                hallazgos.append([f'<b>{_esc_html(h.get("titulo", ""))}</b>', _esc_html(h.get("descripcion", ""))])
        if hallazgos:
            out += _titulo_seccion.format("Hallazgos") + _tabla(["Hallazgo", "Detalle"], hallazgos, ["30%", "70%"])
        senales = []
        for s in (r.get("senales") or []):
            if isinstance(s, str): s = {"titulo": s, "descripcion": "", "severidad": "media"}
            if isinstance(s, dict):
                senales.append([_sev_celda(s.get("severidad", "")),
                                f'<b>{_esc_html(s.get("titulo", ""))}</b><br/><span style="color:#5F6877;">{_esc_html(s.get("descripcion", ""))}</span>'])
        if senales:
            out += _titulo_seccion.format("Señales detectadas") + _tabla(["Severidad", "Señal"], senales, ["13%", "87%"])
        for titulo, items, campo in (("Omisiones", r.get("omisiones"), "descripcion"),
                                     ("Preguntas probables", r.get("preguntas"), "pregunta"),
                                     ("Recomendaciones", r.get("recomendaciones"), "descripcion")):
            lst = _lista(items, campo)
            if lst:
                out += _titulo_seccion.format(titulo) + lst
        return out

    bloques_agentes = ""
    for key, data in por_agente.items():
        nombre = _esc_html(data["nombre"])
        analisis_html = ""
        for a in data["analisis"]:
            r = a.get("resultado", {})
            if not isinstance(r, dict): r = {}
            analisis_html += render_analisis_agente(r, r.get("score", 0))
        chats_html = ""
        for c in data["chats"]:
            pregunta_limpia = _limpiar_unicode(str(c.get("pregunta", "")))
            chats_html += (
                f'<table width="100%" cellpadding="0" cellspacing="0" style="margin:4pt 0;"><tr>'
                f'<td style="background-color:#F3F5F8;border-left:3pt solid #1B2F4B;padding:5pt 8pt;">'
                f'<b style="color:#1B2F4B;">Consulta:</b> {_esc_html(pregunta_limpia)}'
                f'<div style="font-size:9pt;color:#2B2F36;margin-top:3pt;">{_render_markdown_a_html(str(c.get("respuesta", "")))}</div>'
                f'</td></tr></table>'
            )
        bloques_agentes += (
            f'<h3 style="font-size:11pt;margin:14pt 0 4pt 0;padding-bottom:2pt;border-bottom:0.75pt solid #C9D0DA;">{nombre}</h3>'
            f'{analisis_html}'
            f'{(_titulo_seccion.format("Conversaciones sobre este tema") + chats_html) if chats_html else ""}'
        )

    from modulos_base import CSS_PDF, CLASIFICACION
    sec_criticos = (
        "<h2>2. Puntos críticos detectados</h2>" + bloque_senales_html(top_senales)) if top_senales else ""
    sec_preguntas = (
        "<h2>3. Preguntas probables del directorio</h2>" + bloque_preguntas_html(top_preguntas)) if top_preguntas else ""
    html_completo = f'''<!DOCTYPE html>
<html lang="es">
<head>
<meta http-equiv="Content-Type" content="text/html; charset=UTF-8"/>
<title>Informe Pre-Directorio - {nombre_doc}</title>
<style>
{CSS_PDF}
  ul, ol {{ margin: 6px 0 10px 22px; padding: 0; }}
  li {{ margin-bottom: 4px; color: #2B2F36; }}
</style>
</head>
<body>
<div id="pdf_cab"><table width="100%"><tr><td class="cab"><b>DIRECTA</b> &nbsp;|&nbsp; Gerencia General</td><td class="cab" align="right">Informe de Preparaci&oacute;n Pre-Directorio</td></tr></table></div>
<div id="pdf_pie"><table width="100%"><tr><td class="pie">{CLASIFICACION}</td><td class="pie" align="right">P&aacute;gina <pdf:pagenumber/> de <pdf:pagecount/></td></tr></table></div>

<table class="portada" width="100%"><tr><td>
  <span class="eyebrow">GERENCIA GENERAL &nbsp;|&nbsp; DIRECTA</span><br/>
  <span class="tit">Informe de Preparaci&oacute;n Pre-Directorio</span><br/>
  <span class="subtit">An&aacute;lisis del documento para la reuni&oacute;n de directorio</span>
</td></tr></table><br/>
<table class="meta" width="100%">
  <tr><td class="k">Documento</td><td>{nombre_doc}</td><td class="k">P&aacute;ginas</td><td>{paginas}</td></tr>
  <tr><td class="k">Fecha de emisi&oacute;n</td><td>{fecha}</td><td class="k">Clasificaci&oacute;n</td><td>Confidencial</td></tr>
</table>

<h2>1. Resumen ejecutivo</h2>
<table width="100%" cellpadding="0" cellspacing="0"><tr>
  <td width="150" class="tile" style="border:0.75pt solid #C9D0DA;padding:12pt 6pt;text-align:center;vertical-align:top;background-color:#F8F9FB;">
    <span style="font-size:30pt;font-weight:bold;color:{score_global_color};">{score_global}</span><br/>
    <span style="font-size:7.5pt;color:#5F6877;">&Iacute;NDICE GLOBAL DE PREPARACI&Oacute;N</span><br/>
    <span style="font-size:9pt;font-weight:bold;color:{score_global_color};">{score_global_label}</span><br/>
    <span style="font-size:7.5pt;color:#5F6877;">{n_evaluados} de {n_total} &aacute;reas evaluadas</span>
  </td>
  <td width="14"></td>
  <td style="vertical-align:top;font-size:9.5pt;color:#2B2F36;line-height:1.5;">{_render_markdown_a_html(resumen_general)}</td>
</tr></table>

{sec_criticos}

{sec_preguntas}

<h2>4. An&aacute;lisis detallado por &aacute;rea</h2>
{bloques_agentes if bloques_agentes else "<p><i>Sin an&aacute;lisis disponibles.</i></p>"}

</body>
</html>'''

    return html_completo


# Reintento ante cortes de conexión con Supabase en la lectura de fragmentos de cada ola
seleccionar_chunks_para_grupo = _reintentar_db(seleccionar_chunks_para_grupo)